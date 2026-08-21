"""ICP 类型驱动文档路由（doc_type 路径，替代知识库概念）。

首页 7 类文档卡片 → /document-types/{docType}/documents 文档管理 → 上传/解析/查看。
查看页四视图：原文 / 向量化 / PG 结构 / 图谱（按 data_form 组合）。
"""
import asyncio
import json
import os
import uuid
import logging
from typing import List, Optional

from fastapi import APIRouter, UploadFile, File, Query, HTTPException, Body

from backend.core.parsers import ParserRegistry
from backend.core.task_manager import task_manager
from backend.services.document_service import parse_icp_document_background
from backend.services.kg_service import generate_icp_kg_background
from backend.services.model_service import (
    resolve_llm_model_config, require_embedding_model_configured,
)
from backend.storage.es_client import es_client
from backend.storage.hugegraph_client import hugegraph_client
from backend.storage.minio_client import minio_client
from backend.storage.postgres_client import postgres_client

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/document-types", tags=["ICP 文档类型"])

_TABLE_TYPES = ("hmc", "fmeca", "mfl")
_IMAGE_TYPES = ("fta", "rbd")
_TEXT_TYPES = ("fault_history", "function_locator")


# ===================== 类型注册表 =====================

@router.get("")
async def list_document_types():
    return postgres_client.list_document_types()


@router.get("/{doc_type}")
async def get_document_type(doc_type: str):
    t = postgres_client.get_document_type(doc_type)
    if not t:
        raise HTTPException(status_code=404, detail=f"文档类型不存在: {doc_type}")
    return t


# ===================== 文档管理 =====================

@router.post("/{doc_type}/documents/upload")
async def upload_icp_documents(doc_type: str, files: List[UploadFile] = File(...)):
    t = postgres_client.get_document_type(doc_type)
    if not t:
        raise HTTPException(status_code=404, detail=f"文档类型不存在: {doc_type}")

    allowed = [fmt.strip().lower() for fmt in (t.get("input_formats") or "").split(",") if fmt.strip()]
    results = []
    for file in files:
        ext = os.path.splitext(file.filename or "")[1].lower()
        if allowed and ext not in allowed:
            raise HTTPException(
                status_code=400,
                detail=f"文件 {file.filename} 格式不支持（{t['name']} 支持 {t['input_formats']}）")

        doc_id = str(uuid.uuid4())
        minio_path = f"icp/{doc_type}/{doc_id}/{file.filename}"
        content = await file.read()
        with open(f"/tmp/{doc_id}_{file.filename}", "wb") as f:
            f.write(content)
        minio_client.upload_file(minio_path, content)

        postgres_client.insert_icp_document(doc_id, doc_type, file.filename, minio_path)
        results.append({"doc_id": doc_id, "file_name": file.filename, "status": "uploaded"})

    return results


@router.get("/{doc_type}/documents")
async def list_icp_documents(doc_type: str, status: Optional[str] = None):
    if not postgres_client.get_document_type(doc_type):
        raise HTTPException(status_code=404, detail=f"文档类型不存在: {doc_type}")
    return postgres_client.list_icp_documents(doc_type, status)


def _get_icp_document(doc_type: str, doc_id: str):
    doc = postgres_client.get_icp_document(doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="文档不存在")
    if doc["doc_type"] != doc_type:
        raise HTTPException(status_code=403, detail="文档不属于该类型")
    return doc


@router.post("/{doc_type}/documents/{doc_id}/parse")
async def parse_icp_document(doc_type: str, doc_id: str, request: dict = Body(default={})):
    doc = _get_icp_document(doc_type, doc_id)
    if doc["status"] not in ("uploaded", "failed", "indexed", "kg_pending", "kg_done"):
        raise HTTPException(status_code=400, detail="文档正在处理中，请勿重复触发")

    # 解析预检：Chat(LLM) 模型（图类 vision/表格分析）+ Embedding 模型（向量化）
    model_config = resolve_llm_model_config(request.get("model_id"))
    require_embedding_model_configured()

    postgres_client.update_icp_document_status(doc_id, "parsing")
    task_id = await task_manager.create_task(
        doc_id=doc_id, kb_id=doc_type, file_name=doc["file_name"], task_type="parse")
    asyncio.create_task(parse_icp_document_background(task_id, doc_id, model_config))

    return {"task_id": task_id, "doc_id": doc_id, "doc_type": doc_type, "status": "parsing"}


@router.get("/{doc_type}/documents/{doc_id}/parse-result")
async def get_parse_result(doc_type: str, doc_id: str):
    doc = _get_icp_document(doc_type, doc_id)
    return {
        "doc_id": doc_id,
        "doc_type": doc_type,
        "status": doc["status"],
        "parse_count": doc.get("parse_count", 0),
        "last_result": doc.get("last_result"),
    }


@router.get("/{doc_type}/documents/{doc_id}/original")
async def get_original(doc_type: str, doc_id: str, page: int = 1, size: int = 100):
    """原文内容：docx 标题树/表格、txt 分块、excel 表格预览（不触发 vision/embedding）。"""
    doc = _get_icp_document(doc_type, doc_id)
    try:
        file_data = minio_client.get_file(doc["minio_path"])
        ext = os.path.splitext(doc["file_name"])[1].lower()
        temp_path = f"/tmp/icp_orig_{doc_id}{ext}"
        with open(temp_path, "wb") as f:
            f.write(file_data)
    except Exception:
        temp_path = f"/tmp/icp_orig_{doc_id}"
    return _build_original(temp_path, doc_type, doc["file_name"], page, size)


@router.get("/{doc_type}/documents/{doc_id}/vectors")
async def get_vectors(doc_type: str, doc_id: str, page: int = 1, size: int = 50):
    _get_icp_document(doc_type, doc_id)
    return es_client.get_doc_vectors("default", doc_id, page, size)


@router.get("/{doc_type}/documents/{doc_id}/rows")
async def get_rows(doc_type: str, doc_id: str, page: int = 1, size: int = 50,
                   keyword: str = None, severity: str = None,
                   api_reportable: str = None, platform_test: str = None):
    """PG 真源表行数据（表格类）。支持 keyword 全列检索与列过滤。"""
    doc = _get_icp_document(doc_type, doc_id)
    filters = {}
    if keyword:
        filters["keyword"] = keyword
    if severity:
        filters["severity"] = severity
    if api_reportable in ("true", "false"):
        filters["api_reportable"] = api_reportable
    if platform_test in ("true", "false"):
        filters["platform_test"] = platform_test

    if doc_type == "hmc":
        return postgres_client.get_hmc_rows(doc_id, page, size, filters)
    if doc_type == "fmeca":
        return postgres_client.get_fmea_rows(doc_id, page, size, filters)
    if doc_type == "mfl":
        return postgres_client.get_mfl_rows(doc_id, page, size, filters)
    raise HTTPException(status_code=400, detail="该类型不是表格类文档，无 PG 行数据")


@router.get("/{doc_type}/documents/{doc_id}/fault-trees")
async def get_fault_trees(doc_type: str, doc_id: str):
    """三元组 JSON（图类，document_fault_trees）。"""
    _get_icp_document(doc_type, doc_id)
    trees = postgres_client.get_document_fault_trees(doc_id)
    for t in trees:
        try:
            t["graph"] = json.loads(t["graph"]) if isinstance(t["graph"], str) else t["graph"]
        except Exception:
            pass
    return trees


@router.get("/{doc_type}/documents/{doc_id}/graph-data")
async def get_icp_graph_data(doc_type: str, doc_id: str):
    """ICP 图类文档的知识图谱数据（HugeGraph 中 kb_id=doc_type，按 doc_id 过滤）。"""
    _get_icp_document(doc_type, doc_id)
    try:
        vertices = hugegraph_client.get_vertices_by_kb_with_details(doc_type)
        edges = hugegraph_client.get_edges_by_kb_with_details(doc_type)

        def contains_doc_id(props_doc_id):
            if isinstance(props_doc_id, list):
                return doc_id in props_doc_id
            if isinstance(props_doc_id, str):
                return doc_id in props_doc_id.split(",")
            return False

        fv = [v for v in vertices if contains_doc_id(v.get("properties", {}).get("doc_id"))]
        fe = [e for e in edges if contains_doc_id(e.get("properties", {}).get("doc_id"))]
        return {"doc_id": doc_id, "doc_type": doc_type, "vertices": fv, "edges": fe,
                "total_vertices": len(fv), "total_edges": len(fe)}
    except Exception as e:
        return {"doc_id": doc_id, "doc_type": doc_type, "vertices": [], "edges": [],
                "total_vertices": 0, "total_edges": 0, "error": str(e)}


@router.post("/{doc_type}/documents/{doc_id}/generate_kg")
async def generate_kg(doc_type: str, doc_id: str, request: dict = Body(default={})):
    """图谱生成（仅 fta/rbd，人工确认三元组后触发）。"""
    doc = _get_icp_document(doc_type, doc_id)
    if doc_type not in _IMAGE_TYPES:
        raise HTTPException(status_code=400, detail="仅 FTA/RBD 图类文档支持知识图谱生成")
    if doc["status"] != "kg_pending":
        raise HTTPException(status_code=400, detail="文档需处于 kg_pending 状态（解析确认三元组后）")

    model_config = resolve_llm_model_config(request.get("model_id"))
    postgres_client.update_icp_document_status(doc_id, "generating_kg")
    task_id = await task_manager.create_task(
        doc_id=doc_id, kb_id=doc_type, file_name=doc["file_name"], task_type="generate_kg")
    asyncio.create_task(generate_icp_kg_background(task_id, doc_id, doc_type, doc["file_name"], model_config))

    return {"task_id": task_id, "doc_id": doc_id, "doc_type": doc_type, "status": "generating_kg"}


@router.delete("/{doc_type}/documents/{doc_id}")
async def delete_icp_document(doc_type: str, doc_id: str):
    """删除文档：真源表/从属数据 + ES 向量 + HugeGraph + MinIO + 主表记录。"""
    doc = _get_icp_document(doc_type, doc_id)

    # 1. PG 真源表与从属数据
    postgres_client.delete_source_rows_by_doc(doc_id)
    # 2. ES 向量（按 doc_id）
    es_client.delete_vectors_by_doc("default", doc_id)
    # 3. HugeGraph（图类可能已生成图谱）
    try:
        hugegraph_client.delete_edges_by_doc_id(doc_id)
        hugegraph_client.delete_vertices_by_doc_id(doc_id)
    except Exception as e:
        logger.warning(f"HugeGraph 删除失败（忽略）: {e}")
    # 4. MinIO 文件
    try:
        minio_client.delete_file(doc["minio_path"])
    except Exception as e:
        logger.warning(f"MinIO 删除失败（忽略）: {e}")
    # 5. 主表记录
    postgres_client.delete_icp_document(doc_id)

    return {"status": "success", "doc_id": doc_id, "message": "文档已删除"}


# ===================== 检索 =====================

@router.post("/search")
async def structured_search(request: dict = Body(...)):
    """doc_type 过滤的混合检索（替代 kb_id）。filters 支持三码/severity/api_reportable/platform_test 等。"""
    doc_type = request.get("doc_type", "")
    query = request.get("query", "")
    k = int(request.get("k", 10))
    filters = request.get("filters") or {}
    if not doc_type:
        raise HTTPException(status_code=400, detail="doc_type 必填")
    if not postgres_client.get_document_type(doc_type):
        raise HTTPException(status_code=404, detail=f"文档类型不存在: {doc_type}")

    query_vector = None
    if query:
        from backend.core.embedding import embedding_service
        query_vector = embedding_service.embed(query)

    hits = es_client.search_structured("default", doc_type, query, query_vector, k, filters)
    return {"doc_type": doc_type, "query": query, "total": len(hits), "hits": hits}


# ===================== 原文提取（不触发 LLM/embedding） =====================

def _build_original(temp_path: str, doc_type: str, file_name: str, page: int = 1, size: int = 100):
    ext = os.path.splitext(file_name)[1].lower()
    blocks = []
    if ext == ".docx":
        from docx import Document
        from backend.core.docx_parser import DocxParser
        doc = Document(temp_path)
        parser = DocxParser(temp_path)
        parser.doc = doc
        cur_block = None
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            is_heading, level, _ = parser.is_heading_paragraph(para)
            if is_heading:
                cur_block = {"title": text, "level": level, "text": ""}
                blocks.append(cur_block)
            else:
                if cur_block is not None:
                    cur_block["text"] = (cur_block["text"] + "\n" + text).strip()
                else:
                    blocks.append({"title": "", "level": 0, "text": text})
        for idx, table in enumerate(doc.tables, start=1):
            rows = [[cell.text.strip().replace("\n", " ") for cell in r.cells] for r in table.rows]
            blocks.append({"title": f"表格 {idx}", "level": 0,
                           "text": " | ".join(rows[0]) if rows else "", "table": rows[:20]})
    elif ext == ".txt":
        from backend.core.docx_parser import TextParser
        tp = TextParser(temp_path)
        if tp.load():
            for key, b in tp.parse_markdown_headers().items():
                blocks.append({"title": b["title"], "level": b.get("level", 1), "text": b["text"]})
    else:
        import pandas as pd
        try:
            xls = pd.read_excel(temp_path, sheet_name=None, header=None, dtype=str, nrows=50)
            for sheet, df in xls.items():
                data = [[("" if pd.isna(c) else str(c)) for c in row] for row in df.values.tolist()]
                blocks.append({"title": f"工作表: {sheet}", "level": 0,
                               "text": " | ".join(data[0]) if data else "", "table": data[:20]})
        except Exception as e:
            blocks.append({"title": "读取失败", "level": 0, "text": str(e)})

    total = len(blocks)
    start = (page - 1) * size
    return {"doc_type": doc_type, "file_name": file_name, "total": total,
            "page": page, "size": size, "blocks": blocks[start:start + size]}
