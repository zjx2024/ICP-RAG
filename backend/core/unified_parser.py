import os
import re
import json
import uuid
import chardet
from io import BytesIO
from typing import Dict, List, Any
from bs4 import BeautifulSoup, NavigableString, Tag, Comment

from backend.core.parser_utils import find_codec, get_text


class TxtParser:
    def __call__(self, fnm, binary=None, chunk_token_num=128, delimiter="\n!?;。；！？"):
        txt = get_text(fnm, binary)
        return self.parser_txt(txt, chunk_token_num, delimiter)

    @classmethod
    def parser_txt(cls, txt, chunk_token_num=128, delimiter="\n!?;。；！？"):
        if not isinstance(txt, str):
            raise TypeError("txt type should be str!")
        cks = [""]
        delimiter = delimiter.encode('utf-8').decode('unicode_escape').encode('latin1').decode('utf-8')

        def add_chunk(t):
            nonlocal cks
            if len(cks[-1]) > chunk_token_num:
                cks.append(t)
            else:
                cks[-1] += t

        dels = []
        s = 0
        for m in re.finditer(r"`([^`]+)`", delimiter, re.I):
            f, t = m.span()
            dels.append(m.group(1))
            dels.extend(list(delimiter[s: f]))
            s = t
        if s < len(delimiter):
            dels.extend(list(delimiter[s:]))
        dels = [re.escape(d) for d in dels if d]
        dels = [d for d in dels if d]
        dels = "|".join(dels)
        secs = re.split(r"(%s)" % dels, txt)
        for sec in secs:
            if re.match(f"^{dels}$", sec):
                continue
            add_chunk(sec)

        return [[c, ""] for c in cks]


class MarkdownParser:
    def __init__(self, chunk_token_num=128):
        self.chunk_token_num = int(chunk_token_num)

    def __call__(self, fnm, binary=None):
        if binary:
            encoding = find_codec(binary)
            txt = binary.decode(encoding, errors="ignore")
        else:
            with open(fnm, "r") as f:
                txt = f.read()
        return self.parse_markdown(txt)

    def parse_markdown(self, markdown_text):
        extractor = MarkdownElementExtractor(markdown_text)
        return extractor.extract_elements()

    def extract_tables_and_remainder(self, markdown_text, separate_tables=True):
        tables = []
        working_text = markdown_text

        def replace_tables_with_rendered_html(pattern, table_list, render=True):
            new_text = ""
            last_end = 0
            for match in pattern.finditer(working_text):
                raw_table = match.group()
                table_list.append(raw_table)
                if separate_tables:
                    new_text += working_text[last_end: match.start()] + "\n\n"
                else:
                    from markdown import markdown
                    html_table = markdown(raw_table, extensions=["markdown.extensions.tables"]) if render else raw_table
                    new_text += working_text[last_end: match.start()] + html_table + "\n\n"
                last_end = match.end()
            new_text += working_text[last_end:]
            return new_text

        if "|" in markdown_text:
            border_table_pattern = re.compile(
                r"""
                (?:\n|^)
                (?:\|.*?\|.*?\|.*?\n)
                (?:\|(?:\s*[:-]+[-| :]*\s*)\|.*?\n)
                (?:\|.*?\|.*?\|.*?\n)+
            """,
                re.VERBOSE,
            )
            working_text = replace_tables_with_rendered_html(border_table_pattern, tables)

            no_border_table_pattern = re.compile(
                r"""
                (?:\n|^)
                (?:\S.*?\|.*?\n)
                (?:(?:\s*[:-]+[-| :]*\s*).*?\n)
                (?:\S.*?\|.*?\n)+
                """,
                re.VERBOSE,
            )
            working_text = replace_tables_with_rendered_html(no_border_table_pattern, tables)

        if "<table>" in working_text.lower():
            html_table_pattern = re.compile(
                r"""
            (?:\n|^)
            \s*
            (?:
                (?:<html[^>]*>\s*<body[^>]*>\s*<table[^>]*>.*?</table>\s*</body>\s*</html>)
                |
                (?:<body[^>]*>\s*<table[^>]*>.*?</table>\s*</body>)
                |
                (?:<table[^>]*>.*?</table>)
            )
            \s*
            (?=\n|$)
            """,
                re.VERBOSE | re.DOTALL | re.IGNORECASE,
            )

            def replace_html_tables():
                nonlocal working_text
                new_text = ""
                last_end = 0
                for match in html_table_pattern.finditer(working_text):
                    raw_table = match.group()
                    tables.append(raw_table)
                    if separate_tables:
                        new_text += working_text[last_end: match.start()] + "\n\n"
                    else:
                        new_text += working_text[last_end: match.start()] + raw_table + "\n\n"
                    last_end = match.end()
                new_text += working_text[last_end:]
                working_text = new_text

            replace_html_tables()

        return working_text, tables


class MarkdownElementExtractor:
    def __init__(self, markdown_content):
        self.markdown_content = markdown_content
        self.lines = markdown_content.split("\n")

    def extract_elements(self):
        sections = []
        i = 0
        while i < len(self.lines):
            line = self.lines[i]

            if re.match(r"^#{1,6}\s+.*$", line):
                element = self._extract_header(i)
                sections.append(element["content"])
                i = element["end_line"] + 1
            elif line.strip().startswith("```"):
                element = self._extract_code_block(i)
                sections.append(element["content"])
                i = element["end_line"] + 1
            elif re.match(r"^\s*[-*+]\s+.*$", line) or re.match(r"^\s*\d+\.\s+.*$", line):
                element = self._extract_list_block(i)
                sections.append(element["content"])
                i = element["end_line"] + 1
            elif line.strip().startswith(">"):
                element = self._extract_blockquote(i)
                sections.append(element["content"])
                i = element["end_line"] + 1
            elif line.strip():
                element = self._extract_text_block(i)
                sections.append(element["content"])
                i = element["end_line"] + 1
            else:
                i += 1

        sections = [section for section in sections if section.strip()]
        return sections

    def _extract_header(self, start_pos):
        return {
            "type": "header",
            "content": self.lines[start_pos],
            "start_line": start_pos,
            "end_line": start_pos,
        }

    def _extract_code_block(self, start_pos):
        end_pos = start_pos
        content_lines = [self.lines[start_pos]]

        for i in range(start_pos + 1, len(self.lines)):
            content_lines.append(self.lines[i])
            end_pos = i
            if self.lines[i].strip().startswith("```"):
                break

        return {
            "type": "code_block",
            "content": "\n".join(content_lines),
            "start_line": start_pos,
            "end_line": end_pos,
        }

    def _extract_list_block(self, start_pos):
        end_pos = start_pos
        content_lines = []

        i = start_pos
        while i < len(self.lines):
            line = self.lines[i]
            if (
                re.match(r"^\s*[-*+]\s+.*$", line)
                or re.match(r"^\s*\d+\.\s+.*$", line)
                or (i > start_pos and not line.strip())
                or (i > start_pos and re.match(r"^\s{2,}[-*+]\s+.*$", line))
                or (i > start_pos and re.match(r"^\s{2,}\d+\.\s+.*$", line))
                or (i > start_pos and re.match(r"^\s+\w+.*$", line))
            ):
                content_lines.append(line)
                end_pos = i
                i += 1
            else:
                break

        return {
            "type": "list_block",
            "content": "\n".join(content_lines),
            "start_line": start_pos,
            "end_line": end_pos,
        }

    def _extract_blockquote(self, start_pos):
        end_pos = start_pos
        content_lines = []

        i = start_pos
        while i < len(self.lines):
            line = self.lines[i]
            if line.strip().startswith(">") or (i > start_pos and not line.strip()):
                content_lines.append(line)
                end_pos = i
                i += 1
            else:
                break

        return {
            "type": "blockquote",
            "content": "\n".join(content_lines),
            "start_line": start_pos,
            "end_line": end_pos,
        }

    def _extract_text_block(self, start_pos):
        end_pos = start_pos
        content_lines = [self.lines[start_pos]]

        i = start_pos + 1
        while i < len(self.lines):
            line = self.lines[i]
            if re.match(r"^#{1,6}\s+.*$", line) or line.strip().startswith("```") or re.match(r"^\s*[-*+]\s+.*$", line) or re.match(r"^\s*\d+\.\s+.*$", line) or line.strip().startswith(">"):
                break
            elif not line.strip():
                if i + 1 < len(self.lines) and (
                    re.match(r"^#{1,6}\s+.*$", self.lines[i + 1])
                    or self.lines[i + 1].strip().startswith("```")
                    or re.match(r"^\s*[-*+]\s+.*$", self.lines[i + 1])
                    or re.match(r"^\s*\d+\.\s+.*$", self.lines[i + 1])
                    or self.lines[i + 1].strip().startswith(">")
                ):
                    break
                else:
                    content_lines.append(line)
                    end_pos = i
                    i += 1
            else:
                content_lines.append(line)
                end_pos = i
                i += 1

        return {
            "type": "text_block",
            "content": "\n".join(content_lines),
            "start_line": start_pos,
            "end_line": end_pos,
        }


ILLEGAL_CHARACTERS_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


class ExcelParser:
    @staticmethod
    def _load_excel_to_workbook(file_like_object):
        if isinstance(file_like_object, str):
            with open(file_like_object, 'rb') as f:
                file_like_object = BytesIO(f.read())
        elif isinstance(file_like_object, bytes):
            file_like_object = BytesIO(file_like_object)

        file_like_object.seek(0)
        file_head = file_like_object.read(4)
        file_like_object.seek(0)

        if not (file_head.startswith(b"PK\x03\x04") or file_head.startswith(b"\xd0\xcf\x11\xe0")):
            try:
                file_like_object.seek(0)
                import pandas as pd
                df = pd.read_csv(file_like_object)
                return ExcelParser._dataframe_to_workbook(df)
            except Exception as e_csv:
                raise Exception(f"Failed to parse CSV: {e_csv}")

        try:
            from openpyxl import load_workbook
            return load_workbook(file_like_object, data_only=True)
        except Exception as e:
            try:
                file_like_object.seek(0)
                import pandas as pd
                dfs = pd.read_excel(file_like_object, sheet_name=None)
                return ExcelParser._dataframe_to_workbook(dfs)
            except Exception as e_pandas:
                raise Exception(f"pandas.read_excel error: {e_pandas}, original openpyxl error: {e}")

    @staticmethod
    def _clean_dataframe(df):
        def clean_string(s):
            if isinstance(s, str):
                return ILLEGAL_CHARACTERS_RE.sub(" ", s)
            return s

        return df.apply(lambda col: col.map(clean_string))

    @staticmethod
    def _dataframe_to_workbook(df):
        if isinstance(df, dict) and len(df) > 1:
            return ExcelParser._dataframes_to_workbook(df)

        df = ExcelParser._clean_dataframe(df)
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"

        for col_num, column_name in enumerate(df.columns, 1):
            ws.cell(row=1, column=col_num, value=column_name)

        for row_num, row in enumerate(df.values, 2):
            for col_num, value in enumerate(row, 1):
                ws.cell(row=row_num, column=col_num, value=value)

        return wb

    @staticmethod
    def _dataframes_to_workbook(dfs):
        from openpyxl import Workbook
        wb = Workbook()
        default_sheet = wb.active
        wb.remove(default_sheet)

        for sheet_name, df in dfs.items():
            df = ExcelParser._clean_dataframe(df)
            ws = wb.create_sheet(title=sheet_name)
            for col_num, column_name in enumerate(df.columns, 1):
                ws.cell(row=1, column=col_num, value=column_name)
            for row_num, row in enumerate(df.values, 2):
                for col_num, value in enumerate(row, 1):
                    ws.cell(row=row_num, column=col_num, value=value)
        return wb

    def __call__(self, fnm):
        wb = ExcelParser._load_excel_to_workbook(fnm)

        res = []
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            rows = list(ws.rows)
            if not rows:
                continue
            ti = list(rows[0])
            for r in list(rows[1:]):
                fields = []
                for i, c in enumerate(r):
                    if not c.value:
                        continue
                    t = str(ti[i].value) if i < len(ti) else ""
                    t += ("：" if t else "") + str(c.value)
                    fields.append(t)
                line = "; ".join(fields)
                if sheetname.lower().find("sheet") < 0:
                    line += " ——" + sheetname
                res.append(line)
        return res

    def markdown(self, fnm):
        import pandas as pd
        file_like_object = BytesIO(fnm) if not isinstance(fnm, str) else fnm
        try:
            file_like_object.seek(0)
            df = pd.read_excel(file_like_object)
        except Exception as e:
            file_like_object.seek(0)
            df = pd.read_csv(file_like_object)
        df = df.replace(r"^\s*$", "", regex=True)
        return df.to_markdown(index=False)


class PptParser:
    def __init__(self):
        super().__init__()

    def __get_bulleted_text(self, paragraph):
        is_bulleted = bool(paragraph._p.xpath("./a:pPr/a:buChar")) or bool(paragraph._p.xpath("./a:pPr/a:buAutoNum")) or bool(paragraph._p.xpath("./a:pPr/a:buBlip"))
        if is_bulleted:
            return f"{'  '* paragraph.level}.{paragraph.text}"
        else:
            return paragraph.text

    def __extract(self, shape):
        try:
            if hasattr(shape, 'has_text_frame') and shape.has_text_frame:
                text_frame = shape.text_frame
                texts = []
                for paragraph in text_frame.paragraphs:
                    if paragraph.text.strip():
                        texts.append(self.__get_bulleted_text(paragraph))
                return "\n".join(texts)

            try:
                shape_type = shape.shape_type
            except NotImplementedError:
                if hasattr(shape, 'text'):
                    return shape.text.strip()
                return ""

            if shape_type == 19:
                tb = shape.table
                rows = []
                for i in range(1, len(tb.rows)):
                    rows.append("; ".join([tb.cell(
                        0, j).text + ": " + tb.cell(i, j).text for j in range(len(tb.columns)) if tb.cell(i, j)]))
                return "\n".join(rows)

            if shape_type == 6:
                texts = []
                for p in sorted(shape.shapes, key=lambda x: (x.top // 10, x.left)):
                    t = self.__extract(p)
                    if t:
                        texts.append(t)
                return "\n".join(texts)

            return ""

        except Exception:
            return ""

    def __call__(self, fnm, from_page=0, to_page=10000, callback=None):
        from pptx import Presentation
        ppt = Presentation(fnm) if isinstance(fnm, str) else Presentation(BytesIO(fnm))
        txts = []
        self.total_page = len(ppt.slides)
        for i, slide in enumerate(ppt.slides):
            if i < from_page:
                continue
            if i >= to_page:
                break
            texts = []
            for shape in sorted(
                    slide.shapes, key=lambda x: ((x.top if x.top is not None else 0) // 10, x.left if x.left is not None else 0)):
                try:
                    txt = self.__extract(shape)
                    if txt:
                        texts.append(txt)
                except Exception:
                    pass
            txts.append("\n".join(texts))
        return txts


BLOCK_TAGS = [
    "h1", "h2", "h3", "h4", "h5", "h6",
    "p", "div", "article", "section", "aside",
    "ul", "ol", "li",
    "table", "pre", "code", "blockquote",
    "figure", "figcaption"
]
TITLE_TAGS = {"h1": "#", "h2": "##", "h3": "###", "h4": "####", "h5": "#####", "h6": "######"}


class HtmlParser:
    def __call__(self, fnm, binary=None, chunk_token_num=512):
        if binary:
            encoding = find_codec(binary)
            txt = binary.decode(encoding, errors="ignore")
        else:
            with open(fnm, "r", encoding=self.get_encoding(fnm)) as f:
                txt = f.read()
        return self.parser_txt(txt, chunk_token_num)

    @staticmethod
    def get_encoding(file):
        with open(file, 'rb') as f:
            tmp = chardet.detect(f.read())
            return tmp['encoding'] or 'utf-8'

    @classmethod
    def parser_txt(cls, txt, chunk_token_num):
        if not isinstance(txt, str):
            raise TypeError("txt type should be string!")

        temp_sections = []
        from bs4 import BeautifulSoup, NavigableString, Tag, Comment
        soup = BeautifulSoup(txt, "html.parser")
        for style_tag in soup.find_all(["style", "script"]):
            style_tag.decompose()
        for div_tag in soup.find_all("div"):
            for script_tag in div_tag.find_all("script"):
                script_tag.decompose()
        for tag in soup.find_all(True):
            if 'style' in tag.attrs:
                del tag.attrs['style']
        for comment in soup.find_all(string=lambda text: isinstance(text, Comment)):
            comment.extract()

        cls.read_text_recursively(soup.body, temp_sections, chunk_token_num=chunk_token_num)
        block_txt_list, table_list = cls.merge_block_text(temp_sections)
        sections = cls.chunk_block(block_txt_list, chunk_token_num=chunk_token_num)
        for table in table_list:
            sections.append(table.get("content", ""))
        return sections

    @classmethod
    def read_text_recursively(cls, element, parser_result, chunk_token_num=512, parent_name=None, block_id=None):
        if isinstance(element, NavigableString):
            content = element.strip()

            def is_valid_html(content):
                try:
                    from bs4 import BeautifulSoup
                    soup = BeautifulSoup(content, "html.parser")
                    return bool(soup.find())
                except Exception:
                    return False

            return_info = []
            if content:
                if is_valid_html(content):
                    from bs4 import BeautifulSoup
                    soup = BeautifulSoup(content, "html.parser")
                    child_info = cls.read_text_recursively(soup, parser_result, chunk_token_num, element.name, block_id)
                    parser_result.extend(child_info)
                else:
                    info = {"content": element.strip(), "tag_name": "inner_text", "metadata": {"block_id": block_id}}
                    if parent_name:
                        info["tag_name"] = parent_name
                    return_info.append(info)
            return return_info
        elif isinstance(element, Tag):
            import html
            if str.lower(element.name) == "table":
                table_info_list = []
                table_id = str(uuid.uuid1())
                table_list = [html.unescape(str(element))]
                for t in table_list:
                    table_info_list.append({"content": t, "tag_name": "table",
                                            "metadata": {"table_id": table_id, "index": table_list.index(t)}})
                return table_info_list
            else:
                block_id = None
                if str.lower(element.name) in BLOCK_TAGS:
                    block_id = str(uuid.uuid1())
                for child in element.children:
                    child_info = cls.read_text_recursively(child, parser_result, chunk_token_num, element.name,
                                                           block_id)
                    parser_result.extend(child_info)
        return []

    @classmethod
    def merge_block_text(cls, parser_result):
        block_content = []
        current_content = ""
        table_info_list = []
        lask_block_id = None
        for item in parser_result:
            content = item.get("content")
            tag_name = item.get("tag_name")
            title_flag = tag_name in TITLE_TAGS
            block_id = item.get("metadata", {}).get("block_id")
            if block_id:
                if title_flag:
                    content = f"{TITLE_TAGS[tag_name]} {content}"
                if lask_block_id != block_id:
                    if lask_block_id is not None:
                        block_content.append(current_content)
                    current_content = content
                    lask_block_id = block_id
                else:
                    current_content += (" " if current_content else "") + content
            else:
                if tag_name == "table":
                    table_info_list.append(item)
                else:
                    current_content += (" " if current_content else "" + content)
        if current_content:
            block_content.append(current_content)
        return block_content, table_info_list

    @classmethod
    def chunk_block(cls, block_txt_list, chunk_token_num=512):
        chunks = []
        current_block = ""
        current_token_count = 0

        for block in block_txt_list:
            block_token_count = len(block)
            if block_token_count > chunk_token_num:
                if current_block:
                    chunks.append(current_block)
                start = 0
                while start < len(block):
                    end = start + chunk_token_num
                    chunks.append(block[start:end])
                    start = end
                current_block = ""
                current_token_count = 0
            else:
                if current_token_count + block_token_count <= chunk_token_num:
                    current_block += ("\n" if current_block else "") + block
                    current_token_count += block_token_count
                else:
                    chunks.append(current_block)
                    current_block = block
                    current_token_count = block_token_count

        if current_block:
            chunks.append(current_block)

        return chunks


class JsonParser:
    def __init__(self, max_chunk_size: int = 2000, min_chunk_size: int | None = None):
        super().__init__()
        self.max_chunk_size = max_chunk_size * 2
        self.min_chunk_size = min_chunk_size if min_chunk_size is not None else max(max_chunk_size - 200, 50)

    def __call__(self, binary):
        encoding = find_codec(binary)
        txt = binary.decode(encoding, errors="ignore")

        if self.is_jsonl_format(txt):
            sections = self._parse_jsonl(txt)
        else:
            sections = self._parse_json(txt)
        return sections

    @staticmethod
    def _json_size(data: dict) -> int:
        return len(json.dumps(data, ensure_ascii=False))

    @staticmethod
    def _set_nested_dict(d: dict, path: list[str], value: Any) -> None:
        for key in path[:-1]:
            d = d.setdefault(key, {})
        d[path[-1]] = value

    def _list_to_dict_preprocessing(self, data: Any) -> Any:
        if isinstance(data, dict):
            return {k: self._list_to_dict_preprocessing(v) for k, v in data.items()}
        elif isinstance(data, list):
            return {str(i): self._list_to_dict_preprocessing(item) for i, item in enumerate(data)}
        else:
            return data

    def _json_split(self, data, current_path: list[str] | None, chunks: list[dict] | None) -> list[dict]:
        current_path = current_path or []
        chunks = chunks or [{}]
        if isinstance(data, dict):
            for key, value in data.items():
                new_path = current_path + [key]
                chunk_size = self._json_size(chunks[-1])
                size = self._json_size({key: value})
                remaining = self.max_chunk_size - chunk_size

                if size < remaining:
                    self._set_nested_dict(chunks[-1], new_path, value)
                else:
                    if chunk_size >= self.min_chunk_size:
                        chunks.append({})
                    self._json_split(value, new_path, chunks)
        else:
            self._set_nested_dict(chunks[-1], current_path, data)
        return chunks

    def _parse_json(self, content: str) -> list[str]:
        sections = []
        try:
            json_data = json.loads(content)
            chunks = self._json_split(json_data, None, None)
            sections = [json.dumps(line, ensure_ascii=False) for line in chunks if line]
        except json.JSONDecodeError:
            pass
        return sections

    def _parse_jsonl(self, content: str) -> list[str]:
        lines = content.strip().splitlines()
        all_chunks = []
        for line in lines:
            if not line.strip():
                continue
            try:
                data = json.loads(line)
                chunks = self._json_split(data, None, None)
                all_chunks.extend(json.dumps(chunk, ensure_ascii=False) for chunk in chunks if chunk)
            except json.JSONDecodeError:
                continue
        return all_chunks

    def is_jsonl_format(self, txt: str, sample_limit: int = 10, threshold: float = 0.8) -> bool:
        lines = [line.strip() for line in txt.strip().splitlines() if line.strip()]
        if not lines:
            return False

        try:
            json.loads(txt)
            return False
        except json.JSONDecodeError:
            pass

        sample_limit = min(len(lines), sample_limit)
        sample_lines = lines[:sample_limit]
        valid_lines = sum(1 for line in sample_lines if self._is_valid_json(line))

        if not valid_lines:
            return False

        return (valid_lines / len(sample_lines)) >= threshold

    def _is_valid_json(self, line: str) -> bool:
        try:
            json.loads(line)
            return True
        except json.JSONDecodeError:
            return False


class PdfParser:
    def __init__(self):
        pass

    def __call__(self, fnm):
        return self.parse_pdf(fnm)

    def parse_pdf(self, fnm):
        import pdfplumber
        import os
        import uuid
        import traceback
        from datetime import datetime
        from backend.storage.minio_client import minio_client
        
        results = []
        try:
            with pdfplumber.open(fnm) as pdf:
                for page_idx, page in enumerate(pdf.pages):
                    page_data = {
                        "text": "",
                        "images": [],
                        "tables": []
                    }
                    
                    text = page.extract_text()
                    if text:
                        page_data["text"] = text.strip()
                    
                    tables = page.extract_tables()
                    for tbl_idx, table in enumerate(tables):
                        if table and len(table) > 1:
                            markdown_lines = []
                            max_cols = max(len(row) for row in table if row)
                            if table[0]:
                                while len(table[0]) < max_cols:
                                    table[0].append("")
                                markdown_lines.append("| " + " | ".join(str(cell) if cell else "" for cell in table[0]) + " |")
                                markdown_lines.append("| " + " | ".join(["---"] * max_cols) + " |")
                                
                                for row in table[1:]:
                                    while len(row) < max_cols:
                                        row.append("")
                                    markdown_lines.append("| " + " | ".join(str(cell) if cell else "" for cell in row) + " |")
                            
                            page_data["tables"].append({
                                "index": tbl_idx,
                                "title": f"表格 {tbl_idx + 1}",
                                "rows": table,
                                "markdown": "\n".join(markdown_lines),
                                "meta_data": {"data": table}
                            })
                    
                    images = page.images
                    for img_idx, img in enumerate(images):
                        img_stream = img.get("stream")
                        if img_stream:
                            try:
                                img_data = img_stream.get_data()
                                
                                import io
                                from PIL import Image
                                
                                try:
                                    img_obj = Image.open(io.BytesIO(img_data))
                                except Exception:
                                    try:
                                        img_obj = Image.open(io.BytesIO(img_data), formats=['PNG', 'JPEG', 'GIF', 'BMP'])
                                    except Exception:
                                        try:
                                            img_obj = Image.open(io.BytesIO(img_data), formats=['PDF'])
                                        except Exception:
                                            img_obj = None
                                
                                if img_obj:
                                    if img_obj.mode == 'CMYK':
                                        img_obj = img_obj.convert('RGB')
                                    
                                    png_buffer = io.BytesIO()
                                    img_obj.save(png_buffer, format='PNG')
                                    png_data = png_buffer.getvalue()
                                else:
                                    png_data = img_data
                                
                                safe_title = f"pdf_page{page_idx+1}_img{img_idx}"
                                filename = f"images/pdf_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{str(uuid.uuid4())[:8]}.png"
                                minio_client.upload_file(filename, png_data)
                                
                                temp_local_path = f"/tmp/pdf_img_{uuid.uuid4().hex}.png"
                                with open(temp_local_path, 'wb') as f:
                                    f.write(png_data)
                                
                                page_data["images"].append({
                                    "index": img_idx,
                                    "title": f"图片 {img_idx + 1}",
                                    "type": "normal",
                                    "description": "",
                                    "minio_path": filename,
                                    "local_path": temp_local_path
                                })
                            except Exception as e:
                                print(f"PDF图片提取错误: {e}")
                                print(traceback.format_exc())
                    
                    results.append(page_data)
        except Exception as e:
            print(f"PDF解析错误: {e}")
            print(traceback.format_exc())
            raise
        return results


class UnifiedParser:
    def __init__(self, doc_path: str, doc_id: str):
        self.doc_path = doc_path
        self.doc_id = doc_id
        self.file_ext = os.path.splitext(doc_path)[1].lower()
        self.parser = None
        self._init_parser()

    def _init_parser(self):
        if self.file_ext == '.docx':
            self.parser = 'docx'
        elif self.file_ext == '.txt':
            self.parser = TxtParser()
        elif self.file_ext == '.md':
            self.parser = MarkdownParser()
        elif self.file_ext in ['.xlsx', '.xls', '.csv']:
            self.parser = ExcelParser()
        elif self.file_ext in ['.pptx', '.ppt']:
            self.parser = PptParser()
        elif self.file_ext == '.html':
            self.parser = HtmlParser()
        elif self.file_ext in ['.json', '.jsonl']:
            self.parser = JsonParser()
        elif self.file_ext == '.pdf':
            self.parser = PdfParser()
        else:
            self.parser = TxtParser()

    def _convert_to_blocks(self, raw_data, parser_type: str) -> Dict:
        blocks = {}
        block_counter = 1

        if parser_type == 'txt':
            for idx, item in enumerate(raw_data):
                if isinstance(item, list):
                    text = item[0]
                else:
                    text = item
                if text.strip():
                    safe_title = re.sub(r'[^\w\u4e00-\u9fff]+', '_', f"段落{idx + 1}"[:30])
                    block_key = f"txt_{block_counter:03d}_{safe_title}"
                    blocks[block_key] = {
                        'block_id': block_key,
                        'title': f"段落 {idx + 1}",
                        'level': 1,
                        'text': text.strip(),
                        'images': [],
                        'tables': []
                    }
                    block_counter += 1

        elif parser_type == 'markdown':
            if isinstance(raw_data, list):
                sections = raw_data
            else:
                extractor = MarkdownElementExtractor(raw_data)
                sections = extractor.extract_elements()
            
            for idx, section in enumerate(sections):
                header_match = re.match(r'^(#+)\s+(.*)$', section)
                if header_match:
                    level = len(header_match.group(1))
                    title = header_match.group(2).strip()
                    text = ""
                else:
                    level = 1
                    title = f"内容块 {idx + 1}"
                    text = section

                safe_title = re.sub(r'[^\w\u4e00-\u9fff]+', '_', title[:30])
                block_key = f"md_{block_counter:03d}_{safe_title}"
                blocks[block_key] = {
                    'block_id': block_key,
                    'title': title,
                    'level': level,
                    'text': text,
                    'images': [],
                    'tables': []
                }
                block_counter += 1

        elif parser_type == 'excel':
            for idx, line in enumerate(raw_data):
                if line.strip():
                    safe_title = re.sub(r'[^\w\u4e00-\u9fff]+', '_', f"数据行{idx + 1}"[:30])
                    block_key = f"excel_{block_counter:03d}_{safe_title}"
                    blocks[block_key] = {
                        'block_id': block_key,
                        'title': f"数据行 {idx + 1}",
                        'level': 1,
                        'text': line.strip(),
                        'images': [],
                        'tables': []
                    }
                    block_counter += 1

        elif parser_type == 'ppt':
            for idx, page_text in enumerate(raw_data):
                if page_text.strip():
                    safe_title = re.sub(r'[^\w\u4e00-\u9fff]+', '_', f"幻灯片{idx + 1}"[:30])
                    block_key = f"ppt_{block_counter:03d}_{safe_title}"
                    blocks[block_key] = {
                        'block_id': block_key,
                        'title': f"幻灯片 {idx + 1}",
                        'level': 1,
                        'text': page_text.strip(),
                        'images': [],
                        'tables': []
                    }
                    block_counter += 1

        elif parser_type == 'html':
            for idx, section in enumerate(raw_data):
                if section.strip():
                    safe_title = re.sub(r'[^\w\u4e00-\u9fff]+', '_', f"区块{idx + 1}"[:30])
                    block_key = f"html_{block_counter:03d}_{safe_title}"
                    blocks[block_key] = {
                        'block_id': block_key,
                        'title': f"HTML 区块 {idx + 1}",
                        'level': 1,
                        'text': section.strip(),
                        'images': [],
                        'tables': []
                    }
                    block_counter += 1

        elif parser_type == 'json':
            for idx, json_str in enumerate(raw_data):
                if json_str.strip():
                    safe_title = re.sub(r'[^\w\u4e00-\u9fff]+', '_', f"JSON{idx + 1}"[:30])
                    block_key = f"json_{block_counter:03d}_{safe_title}"
                    blocks[block_key] = {
                        'block_id': block_key,
                        'title': f"JSON 块 {idx + 1}",
                        'level': 1,
                        'text': json_str.strip(),
                        'images': [],
                        'tables': []
                    }
                    block_counter += 1

        elif parser_type == 'pdf':
            for idx, page_data in enumerate(raw_data):
                page_text = page_data.get("text", "") if isinstance(page_data, dict) else str(page_data)
                page_images = page_data.get("images", []) if isinstance(page_data, dict) else []
                page_tables = page_data.get("tables", []) if isinstance(page_data, dict) else []
                
                if page_text.strip() or page_images or page_tables:
                    safe_title = re.sub(r'[^\w\u4e00-\u9fff]+', '_', f"PDF页{idx + 1}"[:30])
                    block_key = f"pdf_{block_counter:03d}_{safe_title}"
                    blocks[block_key] = {
                        'block_id': block_key,
                        'title': f"PDF 页面 {idx + 1}",
                        'level': 1,
                        'text': page_text.strip(),
                        'images': page_images,
                        'tables': page_tables
                    }
                    block_counter += 1

        if not blocks:
            blocks['fallback_001'] = {
                'block_id': 'fallback_001',
                'title': '文档内容',
                'level': 1,
                'text': str(raw_data) if raw_data else "",
                'images': [],
                'tables': []
            }

        return blocks

    def process(self) -> Dict:
        if self.parser == 'docx':
            return {"error": "docx format should use DocxParser"}

        try:
            if self.file_ext in ['.json', '.jsonl']:
                with open(self.doc_path, 'rb') as f:
                    binary_data = f.read()
                raw_data = self.parser(binary_data)
            else:
                raw_data = self.parser(self.doc_path)

            parser_type_map = {
                '.txt': 'txt',
                '.md': 'markdown',
                '.xlsx': 'excel',
                '.xls': 'excel',
                '.csv': 'excel',
                '.pptx': 'ppt',
                '.ppt': 'ppt',
                '.html': 'html',
                '.json': 'json',
                '.jsonl': 'json',
                '.pdf': 'pdf'
            }

            blocks = self._convert_to_blocks(raw_data, parser_type_map.get(self.file_ext, 'txt'))

            return {
                "doc_id": self.doc_id,
                "blocks": blocks,
                "total_blocks": len(blocks)
            }

        except Exception as e:
            return {"error": str(e)}