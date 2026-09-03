"""Ordered, lossless-enough intermediate representation for ICP documents.

The older ``DocumentParser`` groups content by Word heading styles and then
attaches tables and images in a second pass.  That is convenient for display,
but loses the actual paragraph/table order and merges headings that only own a
table.  Domain parsers use this module when source order and row provenance
matter.  It performs no database, model, or ArtifactStore writes.
"""
from __future__ import annotations

import hashlib
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable, List, Optional, Sequence, Tuple, Union

from backend.core.parsers.records import SourceRef, source_file_name


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\r", " ").replace("\n", " ").split())


@dataclass(frozen=True)
class CellIR:
    raw_value: Any
    text: str
    source: SourceRef
    anchor_row: int
    anchor_column: int
    formula: str = ""

    @property
    def is_merge_continuation(self) -> bool:
        return (self.source.row_index, self.source.column_index) != (
            self.anchor_row,
            self.anchor_column,
        )


@dataclass(frozen=True)
class ImageIR:
    relation_id: str
    content_type: str
    original_name: str
    sha256: str
    source: SourceRef


@dataclass
class ParagraphIR:
    text: str
    raw_text: str
    style_name: str
    heading_level: int
    section_path: Tuple[str, ...]
    source: SourceRef
    images: List[ImageIR] = field(default_factory=list)
    kind: str = field(init=False, default="paragraph")


@dataclass
class TableIR:
    rows: List[List[CellIR]]
    section_path: Tuple[str, ...]
    source: SourceRef
    images: List[ImageIR] = field(default_factory=list)
    hidden_rows: Tuple[int, ...] = ()
    hidden_columns: Tuple[int, ...] = ()
    kind: str = field(init=False, default="table")

    @property
    def width(self) -> int:
        return max((len(row) for row in self.rows), default=0)

    def matrix(self, *, raw: bool = False) -> List[List[Any]]:
        return [
            [cell.raw_value if raw else cell.text for cell in row]
            for row in self.rows
        ]


DocumentElement = Union[ParagraphIR, TableIR]


@dataclass
class DocumentIR:
    file_name: str
    file_type: str
    elements: List[DocumentElement]

    @property
    def paragraphs(self) -> List[ParagraphIR]:
        return [item for item in self.elements if isinstance(item, ParagraphIR)]

    @property
    def tables(self) -> List[TableIR]:
        return [item for item in self.elements if isinstance(item, TableIR)]

    @property
    def images(self) -> List[ImageIR]:
        result: List[ImageIR] = []
        for item in self.elements:
            result.extend(item.images)
        return result


_NUMBERED_HEADING_RE = re.compile(
    r"^(?P<number>\d+(?:[.．]\d+){0,5})\s+(?P<title>\S.{0,100})$"
)
_CHAPTER_HEADING_RE = re.compile(r"^第\s*[一二三四五六七八九十百零〇0-9]+\s*[章节篇]\s*\S+")


def infer_heading_level(paragraph: Any) -> int:
    """Infer a heading without requiring one particular Word locale/style."""

    text = clean_text(getattr(paragraph, "text", ""))
    if not text:
        return 0
    style = getattr(paragraph, "style", None)
    style_name = clean_text(getattr(style, "name", ""))
    lowered = style_name.lower()
    match = re.search(r"(?:heading|标题)\s*([1-9])", lowered)
    if match:
        return int(match.group(1))
    if lowered in {"title", "subtitle", "题名", "副标题"}:
        return 1

    # Some templates use custom style names but retain the Word outline level.
    try:
        outline = paragraph._p.pPr.outlineLvl
        if outline is not None and outline.val is not None:
            return min(int(outline.val) + 1, 9)
    except Exception:
        pass

    # Conservative fallback for manually formatted engineering headings.
    if "list" not in lowered and "列表" not in style_name and len(text) <= 110:
        numbered = _NUMBERED_HEADING_RE.match(text)
        if numbered and not text.endswith(("。", "；", ";", "，", ",")):
            return min(numbered.group("number").replace("．", ".").count(".") + 1, 9)
        if _CHAPTER_HEADING_RE.match(text):
            return 1
    return 0


def _section_after_heading(
    current: Sequence[str], heading: str, level: int
) -> Tuple[str, ...]:
    values = list(current[: max(level - 1, 0)])
    while len(values) < level - 1:
        values.append("")
    values.append(clean_text(heading))
    return tuple(values)


def _iter_docx_blocks(document: Any) -> Iterable[Any]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _relationship_ids(element: Any) -> List[str]:
    from docx.oxml.ns import qn

    result: List[str] = []
    try:
        for blip in element.xpath(".//a:blip"):
            rel_id = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
            if rel_id:
                result.append(rel_id)
    except Exception:
        pass
    try:
        for image in element.iter("{urn:schemas-microsoft-com:vml}imagedata"):
            rel_id = image.get(qn("r:id"))
            if rel_id:
                result.append(rel_id)
    except Exception:
        pass
    return result


def _images_for_element(
    element: Any,
    part: Any,
    source: SourceRef,
    image_counter: List[int],
) -> List[ImageIR]:
    images: List[ImageIR] = []
    for rel_id in _relationship_ids(element):
        related = getattr(part, "related_parts", {}).get(rel_id)
        if related is None:
            continue
        blob = getattr(related, "blob", b"") or b""
        if not blob:
            continue
        image_counter[0] += 1
        part_name = str(getattr(related, "partname", "") or "")
        images.append(
            ImageIR(
                relation_id=rel_id,
                content_type=str(getattr(related, "content_type", "") or ""),
                original_name=os.path.basename(part_name) or f"image-{image_counter[0]}",
                sha256=hashlib.sha256(blob).hexdigest(),
                source=source.child(element_type="image", image_index=image_counter[0]),
            )
        )
    return images


def _docx_table_ir(
    table: Any,
    *,
    file_name: str,
    sequence: int,
    table_index: int,
    section_path: Tuple[str, ...],
    image_counter: List[int],
) -> TableIR:
    table_source = SourceRef(
        file_name=file_name,
        element_type="table",
        sequence=sequence,
        section_path=section_path,
        block_id=f"tbl-{table_index}",
        table_index=table_index,
    )
    # Keep the XML objects themselves alive as dictionary keys.  Using
    # ``id(cell._tc)`` is unsafe because python-docx may create short-lived
    # proxy objects while walking the grid; CPython can then reuse an id for an
    # unrelated cell and falsely mark it as a merge continuation.
    anchor_by_xml_cell: dict[Any, Tuple[int, int]] = {}
    rows: List[List[CellIR]] = []
    images: List[ImageIR] = []
    image_cells_seen: set[Any] = set()
    for row_no, row in enumerate(table.rows, start=1):
        cells: List[CellIR] = []
        for column_no, cell in enumerate(row.cells, start=1):
            xml_key = cell._tc
            anchor = anchor_by_xml_cell.setdefault(xml_key, (row_no, column_no))
            raw = cell.text or ""
            cell_source = table_source.child(
                element_type="cell", row_index=row_no, column_index=column_no
            )
            cells.append(
                CellIR(
                    raw_value=raw,
                    text=clean_text(raw),
                    source=cell_source,
                    anchor_row=anchor[0],
                    anchor_column=anchor[1],
                )
            )
            # A merged grid position can expose the same XML cell multiple
            # times through ``row.cells``. Scan its drawings once so one source
            # image does not become several synthetic evidence records.
            if xml_key not in image_cells_seen:
                image_cells_seen.add(xml_key)
                for paragraph in cell.paragraphs:
                    images.extend(
                        _images_for_element(
                            paragraph._p, paragraph.part, cell_source, image_counter
                        )
                    )
        rows.append(cells)
    return TableIR(rows=rows, section_path=section_path, source=table_source, images=images)


def _load_docx_ir(file_path: str, *, source_name: Optional[str] = None) -> DocumentIR:
    from docx import Document
    from docx.table import Table

    document = Document(file_path)
    file_name = source_file_name(file_path, source_name)
    elements: List[DocumentElement] = []
    section_path: Tuple[str, ...] = ()
    table_index = 0
    image_counter = [0]

    for sequence, block in enumerate(_iter_docx_blocks(document), start=1):
        if isinstance(block, Table):
            table_index += 1
            elements.append(
                _docx_table_ir(
                    block,
                    file_name=file_name,
                    sequence=sequence,
                    table_index=table_index,
                    section_path=section_path,
                    image_counter=image_counter,
                )
            )
            continue

        raw_text = block.text or ""
        text = clean_text(raw_text)
        heading_level = infer_heading_level(block)
        if heading_level and text:
            section_path = _section_after_heading(section_path, text, heading_level)
        source = SourceRef(
            file_name=file_name,
            element_type="paragraph",
            sequence=sequence,
            section_path=section_path,
            block_id=f"p-{sequence}",
        )
        style = getattr(block, "style", None)
        elements.append(
            ParagraphIR(
                text=text,
                raw_text=raw_text,
                style_name=clean_text(getattr(style, "name", "")),
                heading_level=heading_level,
                section_path=section_path,
                source=source,
                images=_images_for_element(block._p, block.part, source, image_counter),
            )
        )

    return DocumentIR(file_name=file_name, file_type="docx", elements=elements)


_ZERO_FORMAT_RE = re.compile(r"^0+$")


def _excel_display(value: Any, number_format: str = "") -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    fmt = str(number_format or "").split(";")[0].strip()
    if isinstance(value, (int, float)) and _ZERO_FORMAT_RE.fullmatch(fmt):
        try:
            integer = int(value)
            if float(value) == integer:
                return str(integer).zfill(len(fmt))
        except Exception:
            pass
    return clean_text(value)


def _load_openpyxl_ir(file_path: str, *, source_name: Optional[str] = None) -> DocumentIR:
    from openpyxl import load_workbook

    formula_book = load_workbook(file_path, read_only=False, data_only=False)
    value_book = load_workbook(file_path, read_only=False, data_only=True)
    file_name = source_file_name(file_path, source_name)
    elements: List[DocumentElement] = []
    try:
        sequence = 0
        for table_index, sheet_name in enumerate(formula_book.sheetnames, start=1):
            formula_sheet = formula_book[sheet_name]
            value_sheet = value_book[sheet_name]
            merged_anchor: dict[Tuple[int, int], Tuple[int, int]] = {}
            for merged in formula_sheet.merged_cells.ranges:
                anchor = (merged.min_row, merged.min_col)
                for row_no in range(merged.min_row, merged.max_row + 1):
                    for column_no in range(merged.min_col, merged.max_col + 1):
                        merged_anchor[(row_no, column_no)] = anchor

            last_row = 0
            last_column = 0
            for row in formula_sheet.iter_rows():
                for cell in row:
                    cached = value_sheet.cell(cell.row, cell.column).value
                    if cell.value not in (None, "") or cached not in (None, ""):
                        last_row = max(last_row, cell.row)
                        last_column = max(last_column, cell.column)
            if not last_row or not last_column:
                continue
            if last_row * last_column > 2_000_000:
                raise ValueError(
                    f"工作表 {sheet_name} 使用区域过大: {last_row}x{last_column}"
                )

            sequence += 1
            source = SourceRef(
                file_name=file_name,
                element_type="table",
                sequence=sequence,
                section_path=(sheet_name,),
                block_id=f"sheet-{table_index}",
                sheet_name=sheet_name,
                table_index=table_index,
            )
            rows: List[List[CellIR]] = []
            for row_no in range(1, last_row + 1):
                row_values: List[CellIR] = []
                for column_no in range(1, last_column + 1):
                    anchor = merged_anchor.get((row_no, column_no), (row_no, column_no))
                    formula_cell = formula_sheet.cell(anchor[0], anchor[1])
                    cached_cell = value_sheet.cell(anchor[0], anchor[1])
                    formula = (
                        str(formula_cell.value)
                        if isinstance(formula_cell.value, str) and formula_cell.value.startswith("=")
                        else ""
                    )
                    chosen = cached_cell.value if formula and cached_cell.value is not None else formula_cell.value
                    cell_source = source.child(
                        element_type="cell", row_index=row_no, column_index=column_no
                    )
                    row_values.append(
                        CellIR(
                            raw_value=formula_cell.value,
                            text=_excel_display(chosen, formula_cell.number_format),
                            source=cell_source,
                            anchor_row=anchor[0],
                            anchor_column=anchor[1],
                            formula=formula,
                        )
                    )
                rows.append(row_values)
            hidden_rows = tuple(
                row_no
                for row_no in range(1, last_row + 1)
                if formula_sheet.row_dimensions[row_no].hidden
            )
            hidden_columns = tuple(
                column_no
                for column_no in range(1, last_column + 1)
                if formula_sheet.column_dimensions[
                    formula_sheet.cell(1, column_no).column_letter
                ].hidden
            )
            elements.append(
                TableIR(
                    rows=rows,
                    section_path=(sheet_name,),
                    source=source,
                    hidden_rows=hidden_rows,
                    hidden_columns=hidden_columns,
                )
            )
    finally:
        formula_book.close()
        value_book.close()
    return DocumentIR(file_name=file_name, file_type=Path(file_path).suffix.lower()[1:], elements=elements)


def _load_xls_ir(file_path: str, *, source_name: Optional[str] = None) -> DocumentIR:
    # pandas delegates legacy BIFF parsing to xlrd.  Keep the error explicit so
    # an advertised .xls file never degrades into an empty successful parse.
    import pandas as pd

    try:
        sheets = pd.read_excel(file_path, sheet_name=None, header=None, dtype=object)
    except ImportError as exc:
        raise ValueError("解析 .xls 需要安装开源依赖 xlrd") from exc
    file_name = source_file_name(file_path, source_name)
    elements: List[DocumentElement] = []
    for table_index, (sheet_name, frame) in enumerate(sheets.items(), start=1):
        frame = frame.dropna(axis=0, how="all").dropna(axis=1, how="all")
        if frame.empty:
            continue
        source = SourceRef(
            file_name=file_name,
            element_type="table",
            sequence=table_index,
            section_path=(str(sheet_name),),
            block_id=f"sheet-{table_index}",
            sheet_name=str(sheet_name),
            table_index=table_index,
        )
        rows: List[List[CellIR]] = []
        for physical_row, (_, series) in enumerate(frame.iterrows(), start=1):
            values: List[CellIR] = []
            for column_no, value in enumerate(series.tolist(), start=1):
                if pd.isna(value):
                    value = ""
                cell_source = source.child(
                    element_type="cell", row_index=physical_row, column_index=column_no
                )
                values.append(
                    CellIR(
                        raw_value=value,
                        text=clean_text(value),
                        source=cell_source,
                        anchor_row=physical_row,
                        anchor_column=column_no,
                    )
                )
            rows.append(values)
        elements.append(TableIR(rows=rows, section_path=(str(sheet_name),), source=source))
    return DocumentIR(file_name=file_name, file_type="xls", elements=elements)


def _load_text_ir(file_path: str, *, source_name: Optional[str] = None) -> DocumentIR:
    content = ""
    for encoding in ("utf-8", "gb18030", "gbk"):
        try:
            with open(file_path, "r", encoding=encoding) as handle:
                content = handle.read()
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("文本文件编码无法识别")

    file_name = source_file_name(file_path, source_name)
    elements: List[DocumentElement] = []
    section_path: Tuple[str, ...] = ()
    for sequence, line in enumerate(content.splitlines(), start=1):
        raw = line
        text = clean_text(raw)
        heading_level = 0
        markdown = re.match(r"^(#{1,6})\s+(.+)$", raw.strip())
        if markdown:
            heading_level = len(markdown.group(1))
            text = clean_text(markdown.group(2))
        elif text:
            numbered = _NUMBERED_HEADING_RE.match(text)
            if numbered and len(text) <= 110:
                heading_level = numbered.group("number").replace("．", ".").count(".") + 1
        if heading_level:
            section_path = _section_after_heading(section_path, text, heading_level)
        source = SourceRef(
            file_name=file_name,
            element_type="paragraph",
            sequence=sequence,
            section_path=section_path,
            block_id=f"line-{sequence}",
        )
        elements.append(
            ParagraphIR(
                text=text,
                raw_text=raw,
                style_name="",
                heading_level=heading_level,
                section_path=section_path,
                source=source,
            )
        )
    return DocumentIR(file_name=file_name, file_type=Path(file_path).suffix.lower()[1:], elements=elements)


def load_document_ir(
    file_path: str, *, source_name: Optional[str] = None
) -> DocumentIR:
    """Read a supported source without calling models or mutating external state."""

    extension = Path(file_path).suffix.lower()
    if extension == ".docx":
        return _load_docx_ir(file_path, source_name=source_name)
    if extension in {".xlsx", ".xlsm"}:
        return _load_openpyxl_ir(file_path, source_name=source_name)
    if extension == ".xls":
        return _load_xls_ir(file_path, source_name=source_name)
    if extension in {".txt", ".md"}:
        return _load_text_ir(file_path, source_name=source_name)
    raise ValueError(f"暂不支持构建文档 IR: {extension or '无扩展名'}")


__all__ = [
    "CellIR",
    "DocumentElement",
    "DocumentIR",
    "ImageIR",
    "ParagraphIR",
    "TableIR",
    "clean_text",
    "infer_heading_level",
    "load_document_ir",
]
