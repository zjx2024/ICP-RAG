"""Extract controlled document-level metadata used by every ICP projection."""

from __future__ import annotations

import os
import re
from typing import Any, Dict, Iterable


FIELD_ALIASES = {
    "document_number": ("文档编号", "文件编号", "文号", "资料编号", "DOCUMENTNO", "DOCUMENTNUMBER"),
    "applicable_system": ("适用系统", "适用设备", "适用产品", "适用机型", "系统型号", "SYSTEM"),
    "applicable_configuration": ("适用构型", "适用配置", "构型", "配置基线", "CONFIGURATION"),
    "document_version": ("文档版本", "文件版本", "版本", "版次", "修订版", "REVISION", "VERSION"),
    "data_source": ("数据来源", "资料来源", "来源", "DATASOURCE"),
    "data_status": ("数据状态", "资料状态", "有效状态", "STATUS"),
    "compiled_date": ("编制日期", "发布日期", "生效日期", "日期", "DATE"),
}


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _label(value: Any) -> str:
    return re.sub(r"[\s_\-—–/\\:：()（）\[\]【】]+", "", _clean(value).upper())


FIELD_MAP = {
    _label(alias): field
    for field, aliases in FIELD_ALIASES.items()
    for alias in aliases
}


def _from_pairs(pairs: Iterable[tuple[Any, Any]]) -> Dict[str, str]:
    result: Dict[str, str] = {}
    for raw_key, raw_value in pairs:
        field = FIELD_MAP.get(_label(raw_key))
        value = _clean(raw_value)
        if field and value:
            result.setdefault(field, value)
    return result


def extract_document_metadata(file_path: str) -> Dict[str, str]:
    """Extract the controlled cover fields from DOCX, XLSX, or text files."""
    extension = os.path.splitext(file_path)[1].lower()
    if extension == ".docx":
        from docx import Document

        document = Document(file_path)
        pairs = []
        for table in document.tables[:5]:
            for row in table.rows[:30]:
                values = [cell.text for cell in row.cells]
                for index in range(0, len(values) - 1, 2):
                    pairs.append((values[index], values[index + 1]))
        return _from_pairs(pairs)

    if extension in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            pairs = []
            for row in sheet.iter_rows(min_row=1, max_row=30, max_col=6, values_only=True):
                for index in range(0, len(row) - 1, 2):
                    pairs.append((row[index], row[index + 1]))
            return _from_pairs(pairs)
        finally:
            workbook.close()

    if extension == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_path, on_demand=True)
        try:
            sheet = workbook.sheet_by_index(0)
            pairs = []
            for row_index in range(min(sheet.nrows, 30)):
                values = sheet.row_values(row_index, 0, min(sheet.ncols, 6))
                for index in range(0, len(values) - 1, 2):
                    pairs.append((values[index], values[index + 1]))
            return _from_pairs(pairs)
        finally:
            workbook.release_resources()

    if extension in {".txt", ".md"}:
        lines = []
        for encoding in ("utf-8-sig", "gb18030", "gbk"):
            try:
                with open(file_path, "r", encoding=encoding) as handle:
                    lines = [handle.readline() for _ in range(50)]
                break
            except UnicodeDecodeError:
                continue
        pairs = []
        for line in lines:
            match = re.match(r"\s*([^：:]{2,20})\s*[：:]\s*(.+?)\s*$", line)
            if match:
                pairs.append((match.group(1), match.group(2)))
        return _from_pairs(pairs)

    return {}


def metadata_search_tags(metadata: Dict[str, Any]) -> Dict[str, Any]:
    """Build exact-filter OpenSearch fields while keeping the full source object."""
    cleaned = {key: _clean(value) for key, value in (metadata or {}).items() if _clean(value)}
    tags: Dict[str, Any] = {"document_metadata": cleaned}
    for key, value in cleaned.items():
        tags[f"{key}_kwd"] = value
    return tags


__all__ = ["extract_document_metadata", "metadata_search_tags"]
