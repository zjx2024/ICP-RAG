import re
import os
import uuid
import json
from typing import Dict, List, Optional, Any
from backend.config.config import config

class Chapter:
    """章节数据结构"""
    def __init__(self, title: str, text: str, tables: List[dict] = None, images: List[dict] = None,
                 start_idx: int = 0, end_idx: int = 0):
        self.title = title
        self.text = text
        self.tables = tables if tables else []
        self.images = images if images else []
        self.start_idx = start_idx
        self.end_idx = end_idx

    def get_full_content(self) -> str:
        """获取章节完整内容（文本 + 表格 + 图片）"""
        content_parts = [self.text]
        
        for table in self.tables:
            table_text = self._format_table(table)
            if table_text:
                content_parts.append(f"\n【表格】{table_text}")
        
        for img in self.images:
            if img.get("type") == "fault_tree":
                content_parts.append(f"\n【故障树】{json.dumps(img.get('content', {}), ensure_ascii=False)}")
            elif img.get("description"):
                content_parts.append(f"\n【图片描述】{img.get('description')}")
        
        return "\n".join(content_parts)

    def _format_table(self, table: dict) -> str:
        """格式化表格内容"""
        if isinstance(table, dict):
            if "rows" in table:
                return "\n".join([" | ".join(str(cell) for cell in row) for row in table["rows"]])
            elif "markdown" in table:
                return table["markdown"]
            elif "meta_data" in table and isinstance(table["meta_data"], dict):
                data = table["meta_data"].get("data", [])
                if data:
                    return "\n".join([" | ".join(str(cell) for cell in row) for row in data])
        return str(table)

    def to_dict(self) -> dict:
        """转换为字典"""
        return {
            "title": self.title,
            "text": self.text,
            "tables": self.tables,
            "images": self.images,
            "start_idx": self.start_idx,
            "end_idx": self.end_idx
        }

class DocumentParser:
    """文档解析器，支持章节分块、表格和图片提取"""
    
    def __init__(self, chunk_size: int = None, chunk_overlap: int = None):
        self.chunk_size = chunk_size or config.CHUNK_SIZE
        self.chunk_overlap = chunk_overlap or config.CHUNK_OVERLAP
    
    def parse(self, file_path: str) -> Dict[str, Any]:
        """解析文档，返回结构化数据"""
        text_content = self._read_file(file_path)
        tables = []
        images = []
        
        file_ext = os.path.splitext(file_path)[1].lower()
        if file_ext == ".docx":
            tables, images = self._extract_from_docx(file_path)
        elif file_ext in [".md", ".txt"]:
            tables = self._extract_markdown_tables(text_content)
        
        chapters = self._split_into_chapters(text_content, tables, images)
        
        return {
            "file_name": os.path.basename(file_path),
            "text_content": text_content,
            "tables": tables,
            "images": images,
            "chapters": [ch.to_dict() for ch in chapters],
            "doc_id": str(uuid.uuid4()),
            "total_chapters": len(chapters),
            "total_tables": len(tables),
            "total_images": len(images)
        }
    
    def _read_file(self, file_path: str) -> str:
        """读取文件内容"""
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="gbk") as f:
                return f.read()
    
    def _extract_from_docx(self, file_path: str) -> tuple:
        """从docx文件中提取表格和图片"""
        tables = []
        images = []
        try:
            from docx import Document
            doc = Document(file_path)
            
            for idx, table in enumerate(doc.tables):
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)
                
                if any(cell for row in table_data for cell in row):
                    tables.append({
                        "index": idx,
                        "title": f"表格 {idx + 1}",
                        "rows": table_data,
                        "markdown": self._table_to_markdown(table_data),
                        "meta_data": {"data": table_data}
                    })
            
            for idx, rel in enumerate(doc.part.rels.values()):
                if "image" in rel.target_ref:
                    images.append({
                        "index": idx,
                        "title": f"图片 {idx + 1}",
                        "type": "normal",
                        "description": ""
                    })
        except ImportError:
            pass
        
        return tables, images
    
    def _table_to_markdown(self, table_data: List[List[str]]) -> str:
        """将表格数据转换为markdown格式"""
        if not table_data:
            return ""
        
        markdown_lines = []
        markdown_lines.append("| " + " | ".join(table_data[0]) + " |")
        markdown_lines.append("| " + " | ".join(["---"] * len(table_data[0])) + " |")
        
        for row in table_data[1:]:
            markdown_lines.append("| " + " | ".join(row) + " |")
        
        return "\n".join(markdown_lines)
    
    def _extract_markdown_tables(self, text: str) -> List[dict]:
        """从markdown文本中提取表格"""
        tables = []
        table_pattern = re.compile(r'(\|.*\|[\r\n]+\|[-:|]+\|[\r\n]+(\|.*\|[\r\n]*)+)', re.DOTALL)
        matches = table_pattern.finditer(text)
        
        for idx, match in enumerate(matches):
            table_text = match.group(0).strip()
            rows = [row.strip() for row in table_text.split("\n") if row.strip()]
            
            table_data = []
            for row in rows:
                if not row.startswith("|"):
                    continue
                cells = [cell.strip() for cell in row.split("|")[1:-1]]
                if cells and not all(cell == "---" for cell in cells):
                    table_data.append(cells)
            
            if table_data:
                tables.append({
                    "index": idx,
                    "title": f"表格 {idx + 1}",
                    "rows": table_data,
                    "markdown": table_text,
                    "meta_data": {"data": table_data}
                })
        
        return tables
    
    def _split_into_chapters(self, text: str, tables: List[dict], images: List[dict]) -> List[Chapter]:
        """按章节边界分块，每个chapter包含：原文段落 + 表格解析 + 图片解析"""
        chapters = []
        
        chapter_pattern = re.compile(
            r'(第[一二三四五六七八九十]+[章节篇部卷])|'
            r'(\d+\.\d+[\s　]+.+)|\d+\.[\s　]+.+|'
            r'(【.+】)|([一二三四五六七八九十]+、.+)',
            re.UNICODE
        )
        
        matches = list(chapter_pattern.finditer(text))
        
        if not matches:
            chunks = self._split_by_size(text)
            for idx, chunk in enumerate(chunks):
                chapters.append(Chapter(
                    title=f"内容片段 {idx + 1}",
                    text=chunk,
                    tables=tables,
                    images=images,
                    start_idx=0,
                    end_idx=len(text)
                ))
            return chapters
        
        start_idx = 0
        for i, match in enumerate(matches):
            chapter_title = match.group(0).strip()
            chapter_start = match.end()
            chapter_end = matches[i+1].start() if i+1 < len(matches) else len(text)
            
            chapter_text = text[chapter_start:chapter_end].strip()
            
            chapter_tables = self._get_items_in_range(tables, chapter_start, chapter_end)
            chapter_images = self._get_items_in_range(images, chapter_start, chapter_end)
            
            if chapter_text.strip() or chapter_tables or chapter_images:
                chapters.append(Chapter(
                    title=chapter_title,
                    text=chapter_text,
                    tables=chapter_tables,
                    images=chapter_images,
                    start_idx=chapter_start,
                    end_idx=chapter_end
                ))
            
            start_idx = chapter_end
        
        if start_idx < len(text) and text[start_idx:].strip():
            remaining_text = text[start_idx:].strip()
            chapters.append(Chapter(
                title="附录",
                text=remaining_text,
                tables=[],
                images=[],
                start_idx=start_idx,
                end_idx=len(text)
            ))
        
        return chapters
    
    def _get_items_in_range(self, items: List[dict], start_idx: int, end_idx: int) -> List[dict]:
        """获取在指定范围内的表格或图片"""
        result = []
        for item in items:
            item_start = item.get("start_idx", 0)
            if start_idx <= item_start < end_idx:
                result.append(item)
        return result
    
    def _split_by_size(self, text: str) -> List[str]:
        """按大小分块"""
        chunks = []
        paragraphs = re.split(r'\n\n+', text.strip())
        
        current_chunk = ""
        for paragraph in paragraphs:
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            
            if len(current_chunk) + len(paragraph) <= self.chunk_size:
                current_chunk += paragraph + "\n\n"
            else:
                if current_chunk:
                    chunks.append(current_chunk.strip())
                
                if len(paragraph) > self.chunk_size:
                    sub_chunks = [paragraph[i:i+self.chunk_size] 
                                for i in range(0, len(paragraph), self.chunk_size - self.chunk_overlap)]
                    chunks.extend(sub_chunks)
                else:
                    current_chunk = paragraph + "\n\n"
        
        if current_chunk:
            chunks.append(current_chunk.strip())
        
        return chunks
    
    def process(self, file_path: str) -> Dict[str, Any]:
        """处理文档，返回分块结果（兼容性接口）"""
        result = self.parse(file_path)
        
        chunks = []
        for idx, chapter in enumerate(result["chapters"]):
            chapter_obj = Chapter(**chapter)
            full_content = chapter_obj.get_full_content()
            
            if len(full_content) > self.chunk_size:
                sub_chunks = self._split_by_size(full_content)
                for sub_idx, sub_chunk in enumerate(sub_chunks):
                    chunks.append({
                        "chunk_id": f"{result['doc_id']}_ch{idx}_sub{sub_idx}",
                        "title": f"{chapter['title']} - 子片段 {sub_idx + 1}",
                        "section": chapter['title'],
                        "content": sub_chunk,
                        "content_type": "text",
                        "chunk_index": len(chunks),
                        "tables": chapter.get('tables', []),
                        "images": chapter.get('images', [])
                    })
            else:
                chunks.append({
                    "chunk_id": f"{result['doc_id']}_ch{idx}",
                    "title": chapter['title'],
                    "section": chapter['title'],
                    "content": full_content,
                    "content_type": "text",
                    "chunk_index": len(chunks),
                    "tables": chapter.get('tables', []),
                    "images": chapter.get('images', [])
                })
        
        return {
            "file_name": result["file_name"],
            "doc_id": result["doc_id"],
            "chunks": chunks,
            "total_chunks": len(chunks),
            "total_chapters": result["total_chapters"],
            "total_tables": result["total_tables"],
            "total_images": result["total_images"],
            "chapters": result["chapters"]
        }