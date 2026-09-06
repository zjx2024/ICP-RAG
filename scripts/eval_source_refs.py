"""Resolve frozen, reviewable source excerpts for offline eval gold.

These are evaluator-only references, never retrieved evidence supplied to the
agent under test. Source changes require deliberate gold review, not automatic
regeneration of expected answers from the current implementation.
"""
from __future__ import annotations

import hashlib
import json
from functools import lru_cache
from pathlib import Path
from typing import Any, Mapping, Sequence


ROOT = Path(__file__).resolve().parents[1]


@lru_cache(maxsize=16)
def _document(path: str, fingerprint: str):
    from docx import Document

    return Document(path)


def resolve_gold_sources(
    sources: Sequence[Mapping[str, Any] | str], *, source_index: str = ""
) -> list[dict[str, Any]]:
    """Validate fingerprints and coordinates; never silently truncate gold."""
    if source_index:
        index_path = (ROOT / source_index).resolve()
        if not index_path.is_relative_to(ROOT):
            raise ValueError("gold source index must be within the repository")
        index = json.loads(index_path.read_text(encoding="utf-8"))["sources"]
        combined: dict[str, Any] = {}
        for key in sources:
            if not isinstance(key, str) or key not in index:
                raise ValueError(f"unknown gold source key: {key}")
            for entry in index[key]:
                path = entry["path"]
                existing = combined.setdefault(path, {**entry, "locators": []})
                if existing["sha256"] != entry["sha256"]:
                    raise ValueError(f"conflicting gold source revisions: {path}")
                for locator in entry["locators"]:
                    if locator not in existing["locators"]:
                        existing["locators"].append(locator)
        sources = list(combined.values())
    resolved = []
    for source in sources:
        if not isinstance(source, Mapping):
            raise ValueError("gold source keys require an explicit source_index")
        path = (ROOT / str(source["path"])).resolve()
        if not path.is_relative_to(ROOT):
            raise ValueError("gold source must be within the repository")
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        if actual != source["sha256"]:
            raise ValueError(f"gold source changed; review required: {source['path']}")
        excerpts = []
        for locator in source["locators"]:
            kind = locator["kind"]
            if kind in {"table_row", "paragraph"}:
                document = _document(str(path), actual)
                if kind == "paragraph":
                    content = document.paragraphs[_index(locator, "index")].text
                else:
                    table = document.tables[_index(locator, "table")]
                    row = table.rows[_index(locator, "row")]
                    headers = table.rows[_index(locator, "header")].cells if "header" in locator else []
                    content = " | ".join(
                        f"{headers[i].text}: {cell.text}" if i < len(headers) else cell.text
                        for i, cell in enumerate(row.cells)
                    )
            elif kind == "lines":
                start, end = _index(locator, "start"), _index(locator, "end")
                lines = path.read_text(encoding="utf-8").splitlines()
                if start > end or end >= len(lines):
                    raise ValueError("invalid gold line range")
                content = "\n".join(lines[start:end + 1])
            else:
                raise ValueError(f"unknown gold locator: {kind}")
            if not content.strip():
                raise ValueError(f"empty gold excerpt: {source['path']} {locator}")
            excerpts.append({"locator": dict(locator), "text": content})
        if not excerpts:
            raise ValueError("gold source requires at least one excerpt")
        resolved.append({"path": source["path"], "sha256": actual, "excerpts": excerpts})
    return resolved


def _index(locator: Mapping[str, Any], key: str) -> int:
    value = locator[key]
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        raise ValueError(f"gold {key} must be a positive 1-based index")
    return value - 1
