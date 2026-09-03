import unittest
import tempfile
from pathlib import Path

from docx import Document

from backend.core.parsers.diagram_extraction import extract_fta_from_ir, extract_rbd_from_ir
from backend.core.parsers.document_ir import load_document_ir
from backend.icp.reliability import calculate_fta, calculate_rbd


ROOT = Path(__file__).resolve().parents[1]


class DiagramExtractionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fta_ir = load_document_ir(str(ROOT / "test-docs" / "04_FTA故障树分析报告.docx"))
        cls.rbd_ir = load_document_ir(str(ROOT / "test-docs" / "05_RBD可靠性框图分析报告.docx"))

    def test_fta_tables_form_eight_traceable_candidates(self):
        models = extract_fta_from_ir(self.fta_ir)
        self.assertEqual(len(models), 8)
        self.assertEqual(sum(len(item.facts["basic_events"]) for item in models), 68)
        # FTA-SYN-004 的事件清单与割集表原先相差一条（BE-08 无割集行，导致
        # 整树按可达性校验被隔离）；已补齐第 8 条割集，总数由 62 变为 63。
        self.assertEqual(sum(len(item.facts["minimal_cut_sets"]) for item in models), 63)
        self.assertTrue(all(item.candidate["source_kind"] == "table_edges" for item in models))
        first = models[0]
        self.assertEqual(first.business_key, "FTA-SYN-001")
        result = calculate_fta(first.candidate["graph"])
        self.assertEqual(result["minimal_cut_set_count"], 6)
        self.assertAlmostEqual(result["top_event_probability"], 3.771e-4, places=7)

    def test_rbd_expressions_form_four_calculable_candidates(self):
        models = extract_rbd_from_ir(self.rbd_ir)
        self.assertEqual(len(models), 4)
        self.assertEqual(sum(len(item.facts["parameters"]) for item in models), 32)
        self.assertEqual(sum(len(item.facts["declared_results"]) for item in models), 16)
        first = models[0]
        result = calculate_rbd(first.candidate["graph"], mission_time=10)
        self.assertEqual(result["success_path_count"], 2)
        self.assertAlmostEqual(result["system_reliability"], 0.99951007, places=7)
        self.assertEqual(first.candidate["source_kind"], "structured_expression")

    def test_sources_are_physical_rows_not_generated_page_numbers(self):
        fta = extract_fta_from_ir(self.fta_ir)[0]
        self.assertIn("table=", fta.facts["basic_events"][0]["source_ref"])
        self.assertIn("row=", fta.facts["basic_events"][0]["source_ref"])
        rbd = extract_rbd_from_ir(self.rbd_ir)[0]
        self.assertIn("row=", rbd.facts["parameters"][0]["source_ref"])

    def test_fta_supports_multilevel_and_repeated_headers(self):
        document = Document()
        document.add_heading("2.1 显示链路 FTA", level=1)
        metadata = document.add_table(rows=2, cols=2)
        metadata.cell(0, 0).text, metadata.cell(0, 1).text = "故障树编号", "FTA-REAL-01"
        metadata.cell(1, 0).text, metadata.cell(1, 1).text = "顶事件", "显示链路失效"
        events = document.add_table(rows=6, cols=4)
        events.cell(0, 0).merge(events.cell(0, 1)).text = "事件定义"
        events.cell(0, 2).text = "定量信息"
        events.cell(0, 3).text = "说明"
        for column, value in enumerate(("事件ID", "事件名称", "事件概率", "主要原因说明")):
            events.cell(1, column).text = value
        for column, value in enumerate(("E-01", "处理模块失效", "0.001", "器件失效")):
            events.cell(2, column).text = value
        for column, value in enumerate(("事件ID", "事件名称", "事件概率", "主要原因说明")):
            events.cell(3, column).text = value
        for column, value in enumerate(("E-02", "接口模块失效", "0.002", "链路中断")):
            events.cell(4, column).text = value
        cuts = document.add_table(rows=3, cols=2)
        cuts.cell(0, 0).text, cuts.cell(0, 1).text = "序号", "最小割集"
        cuts.cell(1, 0).text, cuts.cell(1, 1).text = "1", "处理模块失效"
        cuts.cell(2, 0).text, cuts.cell(2, 1).text = "2", "接口模块失效"

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fta-headers.docx"
            document.save(path)
            models = extract_fta_from_ir(load_document_ir(str(path)))

        self.assertEqual(len(models), 1)
        self.assertEqual(len(models[0].facts["basic_events"]), 2)
        self.assertEqual(len(models[0].facts["minimal_cut_sets"]), 2)
        self.assertTrue(models[0].facts["basic_events"][0]["raw_data"])

    def test_rbd_supports_multilevel_and_repeated_parameter_headers(self):
        document = Document()
        document.add_heading("3.1 处理链路 RBD", level=1)
        document.add_paragraph("结构表达式：SERIES(A, B)")
        document.add_paragraph("系统成功判据：A 与 B 均工作")
        parameters = document.add_table(rows=6, cols=4)
        parameters.cell(0, 0).merge(parameters.cell(0, 1)).text = "模块信息"
        parameters.cell(0, 2).merge(parameters.cell(0, 3)).text = "可靠性参数"
        for column, value in enumerate(("模块ID", "缩写", "失效率λ（1/h）", "可靠度模型")):
            parameters.cell(1, column).text = value
        for column, value in enumerate(("MOD-A", "A", "1e-5", "指数分布")):
            parameters.cell(2, column).text = value
        for column, value in enumerate(("模块ID", "缩写", "失效率λ（1/h）", "可靠度模型")):
            parameters.cell(3, column).text = value
        for column, value in enumerate(("MOD-B", "B", "2e-5", "指数分布")):
            parameters.cell(4, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "rbd-headers.docx"
            document.save(path)
            models = extract_rbd_from_ir(load_document_ir(str(path)))

        self.assertEqual(len(models), 1)
        self.assertEqual(len(models[0].facts["parameters"]), 2)
        self.assertEqual(models[0].facts["parameters"][1]["module_id"], "MOD-B")

    def test_rbd_numbers_fail_closed_and_parameter_evidence_survives_without_topology(self):
        document = Document()
        document.add_heading("4.1 参数证据 RBD", level=1)
        parameters = document.add_table(rows=2, cols=5)
        headers = ["模块ID", "模块名称", "失效率λ（1/h）", "可靠度", "修复率μ（1/h）"]
        values = ["MOD-A", "处理模块", "1.2×10^-5 /h", ">=0.99", "2E-3/h"]
        for column, value in enumerate(headers):
            parameters.cell(0, column).text = value
        for column, value in enumerate(values):
            parameters.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "rbd-parameters-only.docx"
            document.save(path)
            models = extract_rbd_from_ir(load_document_ir(str(path)))

        self.assertEqual(len(models), 1)
        self.assertEqual(models[0].candidate["graph"]["vertices"], [])
        self.assertFalse(models[0].candidate["graph"]["metadata"]["topology_complete"])
        parameter = models[0].facts["parameters"][0]
        self.assertAlmostEqual(parameter["failure_rate"], 1.2e-5)
        self.assertIsNone(parameter["reliability"])
        self.assertEqual(parameter["reliability_raw"], ">=0.99")
        self.assertAlmostEqual(parameter["repair_rate"], 2e-3)
        self.assertTrue(parameter["normalization_warnings"])

    def test_invalid_rbd_expression_still_retains_parameter_candidate(self):
        document = Document()
        document.add_heading("4.2 表达式错误 RBD", level=1)
        document.add_paragraph("结构表达式：SERIES(A")
        parameters = document.add_table(rows=2, cols=3)
        for column, value in enumerate(("模块ID", "缩写", "失效率")):
            parameters.cell(0, column).text = value
        for column, value in enumerate(("MOD-A", "A", "1e-5")):
            parameters.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "rbd-invalid-expression.docx"
            document.save(path)
            models = extract_rbd_from_ir(load_document_ir(str(path)))

        self.assertEqual(len(models), 1)
        self.assertEqual(
            models[0].candidate["graph"]["metadata"]["topology_source"],
            "invalid_expression_parameters",
        )
        self.assertEqual(len(models[0].facts["parameters"]), 1)
        self.assertTrue(any("不得确认或计算" in warning for warning in models[0].warnings))


if __name__ == "__main__":
    unittest.main()
