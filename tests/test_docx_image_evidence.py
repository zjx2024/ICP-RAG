import base64
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import parse_xml

from backend.core import docx_parser as module
from backend.icp.image_analysis import analyze_icp_diagram


class DocxImageEvidenceTests(unittest.TestCase):
    def test_embedded_image_keeps_declared_format_and_stable_content_key(self):
        parser = module.DocxParser("unused.docx")
        jpeg = b"\xff\xd8\xff\xe0mock-jpeg"

        with patch.object(module.artifact_store, "upload_file") as upload:
            first_key, first_url = parser.save_image_artifact(
                jpeg,
                3,
                "h1_总体架构",
                "doc-1",
                content_type="image/jpeg",
                source_name="/word/media/image3.jpeg",
            )
            second_key, second_url = parser.save_image_artifact(
                jpeg,
                3,
                "h1_总体架构",
                "doc-1",
                content_type="image/jpeg",
                source_name="/word/media/image3.jpeg",
            )

        self.assertEqual(first_key, second_key)
        self.assertEqual(first_url, second_url)
        self.assertTrue(first_key.endswith(".jpg"))
        self.assertEqual(f"/{first_key}", first_url)
        self.assertEqual("image/jpeg", upload.call_args.kwargs["content_type"])
        self.assertEqual(first_key.rsplit("/", 1)[-1], upload.call_args.kwargs["original_name"])

    def test_unknown_word_drawing_is_not_falsely_labelled_as_png(self):
        parser = module.DocxParser("unused.docx")
        with patch.object(module.artifact_store, "upload_file") as upload:
            key, _url = parser.save_image_artifact(
                b"unknown-drawing",
                1,
                "架构图",
                "doc-2",
                content_type="application/octet-stream",
                source_name="/word/media/image1.unknown",
            )

        self.assertTrue(key.endswith(".bin"))
        self.assertEqual("application/octet-stream", upload.call_args.kwargs["content_type"])

    def test_caption_detection_does_not_promote_nearby_prose(self):
        parser = module.DocxParser("unused.docx")
        self.assertEqual("系统数据流", parser.extract_title_from_text("图 2-3 系统数据流", "image"))
        self.assertEqual("", parser.extract_title_from_text("系统由三个模块组成", "image"))
        self.assertEqual("参数清单", parser.extract_title_from_text("表 4.1 参数清单", "table"))
        self.assertEqual("", parser.extract_title_from_text("参见下表说明", "table"))

    def test_table_cell_image_is_saved_with_mime_and_source_locator(self):
        png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        document = Document()
        document.add_heading("1 系统架构", level=1)
        document.add_paragraph("架构图如下")
        table = document.add_table(rows=1, cols=1)

        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "source.png"
            image_path.write_bytes(png)
            table.cell(0, 0).paragraphs[0].add_run().add_picture(str(image_path))
            doc_path = Path(directory) / "table-image.docx"
            document.save(doc_path)
            with patch.object(module.artifact_store, "upload_file") as upload:
                parsed = module.DocumentParser(str(doc_path), "doc-table-image").process()

        images = [
            image
            for block in parsed["blocks"].values()
            for image in block["images"]
        ]
        self.assertEqual(len(images), 1)
        self.assertEqual(images[0]["content_type"], "image/png")
        self.assertEqual(images[0]["container"], "table")
        self.assertIn("container=table", images[0]["source_ref"])
        self.assertEqual(upload.call_args.kwargs["content_type"], "image/png")

    def test_vml_and_textbox_relationships_are_discovered(self):
        element = parse_xml(
            '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'xmlns:v="urn:schemas-microsoft-com:vml">'
            '<w:r><w:drawing><w:txbxContent><a:blip r:embed="rIdTextbox"/>'
            '</w:txbxContent></w:drawing></w:r>'
            '<w:r><w:pict><v:shape><v:imagedata r:id="rIdVml"/>'
            '</v:shape></w:pict></w:r></w:p>'
        )
        self.assertEqual(
            module.DocumentParser._image_relationship_ids(element),
            ["rIdTextbox", "rIdVml"],
        )

    def test_unsupported_engineering_mime_skips_vision_without_data_uri(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "drawing.tiff"
            path.write_bytes(b"II*\x00mock-tiff")
            with patch(
                "backend.icp.image_analysis.llm_service._call_model"
            ) as call_model, patch(
                "backend.icp.image_analysis.llm_service.encode_image_to_base64"
            ) as encode:
                result = analyze_icp_diagram(str(path), "fta", "工程图")

        self.assertTrue(result["metadata"]["vision_skipped"])
        self.assertEqual(result["metadata"]["mime_type"], "image/tiff")
        encode.assert_not_called()
        call_model.assert_not_called()

    def test_diagram_vision_uses_bounded_non_reasoning_json_config(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "drawing.png"
            path.write_bytes(b"\x89PNG\r\n\x1a\nmock")
            with patch(
                "backend.icp.image_analysis.llm_service.get_image_mime_type",
                return_value="image/png",
            ), patch(
                "backend.icp.image_analysis.llm_service.encode_image_to_base64",
                return_value="encoded",
            ), patch(
                "backend.icp.image_analysis.llm_service._call_model",
                return_value='{"vertices": [], "edges": []}',
            ) as call_model:
                analyze_icp_diagram(
                    str(path),
                    "fta",
                    "工程图",
                    model_config={"provider": "ollama", "timeout": 3000},
                )

        effective = call_model.call_args.args[1]
        self.assertEqual(effective["num_predict"], 2048)
        self.assertFalse(effective["reasoning"])
        self.assertEqual(effective["format"], "json")
        self.assertFalse(effective["stream"])
        self.assertEqual(effective["timeout"], 120.0)


if __name__ == "__main__":
    unittest.main()
