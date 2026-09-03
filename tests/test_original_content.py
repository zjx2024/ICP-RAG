import base64
import os
import tempfile
import unittest
from io import BytesIO

from backend.core.original_content import (
    build_evidence_preview,
    build_original_content,
    read_original_docx_image,
)


_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zl9sAAAAASUVORK5CYII="
)


class OriginalContentTests(unittest.TestCase):
    def test_docx_keeps_body_order_and_full_table_and_image(self):
        from docx import Document
        from docx.shared import Inches

        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, "架构说明.docx")
            document = Document()
            document.add_paragraph("表格之前的段落")
            table = document.add_table(rows=26, cols=2)
            for row_index, row in enumerate(table.rows):
                row.cells[0].text = f"行 {row_index + 1}"
                row.cells[1].text = f"值 {row_index + 1}"
            document.add_paragraph("表格之后的段落")
            picture_paragraph = document.add_paragraph()
            picture_paragraph.add_run().add_picture(BytesIO(_TINY_PNG), width=Inches(0.2))
            document.save(path)

            result = build_original_content(
                path,
                doc_type="function_locator",
                doc_id="doc-1",
                file_name="架构说明.docx",
                page=1,
                size=100,
            )

            self.assertEqual(
                [block["kind"] for block in result["blocks"]],
                ["paragraph", "table", "paragraph", "image"],
            )
            self.assertEqual(len(result["blocks"][1]["table"]), 26)
            self.assertEqual(result["counts"], {"paragraphs": 2, "tables": 1, "images": 1})
            self.assertTrue(result["blocks"][3]["image_url"].endswith("/embedded-images/1"))

            blob, content_type, _name = read_original_docx_image(
                path,
                1,
                source_name="架构说明.docx",
            )
            self.assertEqual(blob, _TINY_PNG)
            self.assertEqual(content_type, "image/png")

    def test_pagination_total_is_stable_after_images_are_ordered(self):
        from docx import Document
        from docx.shared import Inches

        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, "FTA.docx")
            document = Document()
            document.add_paragraph("段落一")
            document.add_paragraph("段落二")
            document.add_paragraph().add_run().add_picture(BytesIO(_TINY_PNG), width=Inches(0.2))
            document.save(path)

            first = build_original_content(
                path,
                doc_type="fta",
                doc_id="doc-2",
                file_name="FTA.docx",
                page=1,
                size=2,
            )
            second = build_original_content(
                path,
                doc_type="fta",
                doc_id="doc-2",
                file_name="FTA.docx",
                page=2,
                size=2,
            )

            self.assertEqual(first["total"], 3)
            self.assertEqual(second["total"], 3)
            self.assertEqual([block["kind"] for block in second["blocks"]], ["image"])

    def test_xlsx_returns_every_used_row(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, "HMC.xlsx")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "HMC清单"
            for row_index in range(1, 81):
                sheet.cell(row_index, 1, f"代码-{row_index:03d}")
                sheet.cell(row_index, 2, f"现象-{row_index}")
            workbook.save(path)
            workbook.close()

            result = build_original_content(
                path,
                doc_type="hmc",
                doc_id="doc-3",
                file_name="HMC.xlsx",
                page=1,
                size=100,
            )

            self.assertEqual(result["total"], 1)
            self.assertEqual(result["blocks"][0]["kind"], "table")
            self.assertEqual(len(result["blocks"][0]["table"]), 80)
            self.assertEqual(result["blocks"][0]["table"][-1][0], "代码-080")

    def test_evidence_preview_selects_matching_table_row_nearby_text_and_image(self):
        blocks = [
            {
                "id": "p-30",
                "kind": "paragraph",
                "sequence": 30,
                "title": "4 外部数据通信中断故障树",
                "text": "",
                "source_ref": "报告.docx#4 外部数据通信中断故障树;block=p-30;seq=30",
            },
            {
                "id": "tbl-9",
                "kind": "table",
                "sequence": 31,
                "title": "表格 9",
                "source_ref": "报告.docx#4 外部数据通信中断故障树;block=tbl-9;table=9",
                "table": [
                    ["字段", "内容"],
                    ["故障树编号", "FTA-SYN-003"],
                    ["顶事件", "系统无法与外部设备交换有效业务数据"],
                ],
            },
            {
                "id": "img-3",
                "kind": "image",
                "sequence": 32,
                "title": "图片 3",
                "source_ref": "报告.docx#4 外部数据通信中断故障树;block=p-32;image=3",
                "image_url": "/api/v1/document-types/fta/documents/doc-1/embedded-images/3",
            },
            {
                "id": "p-33",
                "kind": "paragraph",
                "sequence": 33,
                "title": "",
                "text": "图4-1 外部数据通信中断故障树",
                "source_ref": "报告.docx#4 外部数据通信中断故障树;block=p-33;seq=33",
            },
        ]

        result = build_evidence_preview(
            blocks,
            source_ref=(
                "materialized.docx#4 外部数据通信中断故障树;"
                "block=tbl-9;table=9"
            ),
            title="系统无法与外部设备交换有效业务数据",
            structure_summary="FTA 结构：14 个节点、13 条关系。",
        )

        self.assertTrue(result["found"])
        self.assertEqual(result["table"]["headers"], ["字段", "内容"])
        self.assertEqual(
            result["table"]["row"],
            ["顶事件", "系统无法与外部设备交换有效业务数据"],
        )
        self.assertEqual(result["excerpt"], "图4-1 外部数据通信中断故障树")
        self.assertTrue(result["image"]["image_url"].endswith("/embedded-images/3"))
        self.assertIn("14 个节点", result["structure_summary"])

    def test_evidence_preview_respects_explicit_source_row(self):
        result = build_evidence_preview(
            [
                {
                    "id": "tbl-2",
                    "kind": "table",
                    "sequence": 2,
                    "source_ref": "HMC.xlsx#HMC清单;block=tbl-2;table=2",
                    "table": [
                        ["代码", "故障现象"],
                        ["001", "输入欠压"],
                        ["002", "输出过压"],
                    ],
                }
            ],
            source_ref="HMC.xlsx#HMC清单;block=tbl-2;table=2;row=3",
            title="未用于选择",
        )
        self.assertEqual(result["table"]["row"], ["002", "输出过压"])


if __name__ == "__main__":
    unittest.main()
