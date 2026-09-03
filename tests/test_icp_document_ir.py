import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from backend.core.parsers.document_ir import ParagraphIR, TableIR, load_document_ir


class DocumentIRTests(unittest.TestCase):
    def test_preserves_docx_body_order_and_table_only_sections(self):
        document = Document()
        document.add_heading("1 系统说明", level=1)
        document.add_paragraph("说明正文")
        first = document.add_table(rows=2, cols=2)
        first.cell(0, 0).text = "字段"
        first.cell(0, 1).text = "值"
        first.cell(1, 0).text = "模块"
        first.cell(1, 1).text = "GPP"
        document.add_heading("1.1 只有表格的章节", level=2)
        second = document.add_table(rows=1, cols=2)
        second.cell(0, 0).text = "代码"
        second.cell(0, 1).text = "A-001"

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ordered.docx"
            document.save(path)
            ir = load_document_ir(str(path))

        kinds = [element.kind for element in ir.elements if isinstance(element, (ParagraphIR, TableIR))]
        self.assertEqual(kinds, ["paragraph", "paragraph", "table", "paragraph", "table"])
        self.assertEqual(ir.tables[0].section_path[-1], "1 系统说明")
        self.assertEqual(ir.tables[1].section_path[-1], "1.1 只有表格的章节")
        self.assertEqual(ir.tables[1].matrix()[0], ["代码", "A-001"])

    def test_marks_merged_docx_grid_cells_with_one_anchor(self):
        document = Document()
        table = document.add_table(rows=2, cols=3)
        merged = table.cell(0, 0).merge(table.cell(0, 2))
        merged.text = "HMC代码"
        for column, value in enumerate(("LCN", "HAMC", "HSC")):
            table.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "merged.docx"
            document.save(path)
            ir = load_document_ir(str(path))

        first_row = ir.tables[0].rows[0]
        self.assertEqual([cell.text for cell in first_row], ["HMC代码"] * 3)
        self.assertFalse(first_row[0].is_merge_continuation)
        self.assertTrue(first_row[1].is_merge_continuation)
        self.assertEqual((first_row[2].anchor_row, first_row[2].anchor_column), (1, 1))

    def test_vertical_docx_merge_keeps_the_original_anchor(self):
        document = Document()
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "代码"
        table.cell(0, 1).text = "现象"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "HMC-01"
        table.cell(1, 1).text = "现象甲"
        table.cell(2, 1).text = "现象乙"

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "vertical-merge.docx"
            document.save(path)
            ir = load_document_ir(str(path))

        anchor = ir.tables[0].rows[1][0]
        continuation = ir.tables[0].rows[2][0]
        self.assertFalse(anchor.is_merge_continuation)
        self.assertTrue(continuation.is_merge_continuation)
        self.assertEqual(
            (continuation.anchor_row, continuation.anchor_column),
            (anchor.source.row_index, anchor.source.column_index),
        )

    def test_preserves_excel_display_zeros_and_merged_headers_after_long_cover(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "HMC数据"
        for row in range(1, 13):
            sheet.cell(row, 1, f"封面字段{row}")
            sheet.cell(row, 2, f"值{row}")
        sheet.merge_cells(start_row=13, start_column=1, end_row=13, end_column=3)
        sheet.cell(13, 1, "HMC代码")
        sheet.cell(14, 1, "LCN")
        sheet.cell(14, 2, "HAMC")
        sheet.cell(14, 3, "HSC")
        sheet.cell(15, 1, 91810001)
        sheet.cell(15, 2, "**")
        sheet.cell(15, 3, 1)
        sheet.cell(15, 3).number_format = "0000"

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "merged.xlsx"
            workbook.save(path)
            ir = load_document_ir(str(path))

        self.assertEqual(ir.tables[0].matrix()[12][:3], ["HMC代码"] * 3)
        self.assertEqual(ir.tables[0].matrix()[14][2], "0001")
        self.assertEqual(ir.tables[0].source.sheet_name, "HMC数据")


if __name__ == "__main__":
    unittest.main()
