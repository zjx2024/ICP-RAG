import tempfile
import unittest
from pathlib import Path

from docx import Document

from backend.core.document_metadata import extract_document_metadata, metadata_search_tags


class DocumentMetadataTests(unittest.TestCase):
    def test_extracts_controlled_docx_cover_and_builds_search_tags(self):
        document = Document()
        table = document.add_table(rows=3, cols=2)
        values = [
            ("文档编号", "SYN-ICP-DOC-002"),
            ("适用系统", "SYN-ICP-9000"),
            ("文档版本", "V1.0"),
        ]
        for row, values_for_row in zip(table.rows, values):
            row.cells[0].text, row.cells[1].text = values_for_row

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "metadata.docx"
            document.save(path)
            metadata = extract_document_metadata(str(path))

        self.assertEqual(metadata["document_number"], "SYN-ICP-DOC-002")
        tags = metadata_search_tags(metadata)
        self.assertEqual(tags["document_version_kwd"], "V1.0")
        self.assertEqual(tags["document_metadata"], metadata)

    def test_merges_split_cover_tables_and_accepts_common_field_aliases(self):
        document = Document()
        first = document.add_table(rows=1, cols=4)
        first.cell(0, 0).text = "文号"
        first.cell(0, 1).text = "ICP-REAL-001"
        first.cell(0, 2).text = "适用机型"
        first.cell(0, 3).text = "ICP-9000"
        second = document.add_table(rows=1, cols=4)
        second.cell(0, 0).text = "构 型"
        second.cell(0, 1).text = "CONFIG-B"
        second.cell(0, 2).text = "版次"
        second.cell(0, 3).text = "R2"

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "metadata-alias.docx"
            document.save(path)
            metadata = extract_document_metadata(str(path))

        self.assertEqual(metadata["document_number"], "ICP-REAL-001")
        self.assertEqual(metadata["applicable_system"], "ICP-9000")
        self.assertEqual(metadata["applicable_configuration"], "CONFIG-B")
        self.assertEqual(metadata["document_version"], "R2")


if __name__ == "__main__":
    unittest.main()
