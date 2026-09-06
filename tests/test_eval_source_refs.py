import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from docx import Document

from scripts import eval_source_refs as sources
from scripts import evaluate_icp_paraphrases as evaluator


class GoldSourceReferenceTests(unittest.TestCase):
    def test_word_coordinates_and_header_preserve_empty_api(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "source.docx"
            document = Document()
            document.add_paragraph("案例：只有在复测期间未再发生。")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = "是否直接上报", "API"
            table.cell(1, 0).text, table.cell(1, 1).text = "否", "—"
            document.save(path)
            reference = {"path": "source.docx", "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "locators": [
                {"kind": "table_row", "table": 1, "row": 2, "header": 1},
                {"kind": "paragraph", "index": 1},
            ]}
            with patch.object(sources, "ROOT", Path(directory)):
                result = sources.resolve_gold_sources([reference])
                self.assertEqual("是否直接上报: 否 | API: —", result[0]["excerpts"][0]["text"])
                self.assertIn("复测期间", result[0]["excerpts"][1]["text"])
                document.add_paragraph("changed")
                document.save(path)
                with self.assertRaisesRegex(ValueError, "review required"):
                    sources.resolve_gold_sources([reference])

    def test_unknown_source_keys_and_invalid_coordinates_fail_closed(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            text = root / "scope.txt"
            text.write_text("first\nsecond\n", encoding="utf-8")
            reference = {"path": "scope.txt", "sha256": hashlib.sha256(text.read_bytes()).hexdigest(),
                         "locators": [{"kind": "lines", "start": 1, "end": 2}]}
            (root / "index.json").write_text(json.dumps({"sources": {"SCOPE": [reference]}}), encoding="utf-8")
            with patch.object(sources, "ROOT", root):
                self.assertEqual(1, len(sources.resolve_gold_sources(["SCOPE", "SCOPE"], source_index="index.json")))
                with self.assertRaisesRegex(ValueError, "unknown gold"):
                    sources.resolve_gold_sources(["MISSING"], source_index="index.json")
                reference["locators"][0]["start"] = 0
                with self.assertRaisesRegex(ValueError, "positive 1-based"):
                    sources.resolve_gold_sources([reference])

    def test_judge_receives_source_context_separately_from_answer(self):
        context = [{"path": "fixture", "excerpts": [{"text": "code A has no API"}]}]
        class Judge:
            def evaluate(self, **values):
                self.values = values
                return [{"id": "identity", "passed": False, "reason": "wrong API"}]
        judge = Judge()
        _, report = evaluator._evaluate_semantic_criteria(
            [{"id": "identity", "criterion": "use the matching row"}], judge=judge,
            question="its API?", prior_turns=[], answer="borrow code B API", reference_context=context,
        )
        self.assertEqual(context, judge.values["reference_context"])
        self.assertEqual("borrow code B API", judge.values["answer"])
        self.assertFalse(report["criteria"][0]["passed"])
