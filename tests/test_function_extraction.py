import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image

from backend.core.parsers.function_extraction import extract_function_drafts


TEST_DOCS = Path(__file__).resolve().parents[1] / "test-docs"


class FunctionExtractionTests(unittest.TestCase):
    def test_out_of_order_excel_tables_merge_function_detail_and_resolve_interfaces(self):
        workbook = Workbook()
        cover = workbook.active
        cover.title = "封面"
        cover.append(["系统ID", "SYS-DEMO-01"])
        cover.append(["系统名称", "演示综合处理系统"])

        symptom = workbook.create_sheet("失效表现")
        symptom.append(["模块ID", "模块名称", "功能ID", "功能名称", "典型失效表现"])
        symptom.append(["MOD-DEMO-A", "输入模块", "FUNC-DEMO-01", "输入处理", "输入不更新；状态字冻结"])

        interface = workbook.create_sheet("接口")
        interface.append(["接口编号", "发送方", "接收方", "关系类型", "信号名称", "协议"])
        interface.append(["IF-DEMO-01", "MOD-DEMO-A", "MOD-DEMO-B", "数据流", "标准化数据", "LOCAL-BUS"])

        functions = workbook.create_sheet("功能清单")
        for row in range(1, 10):
            functions.cell(row, 1, f"封面说明{row}")
        functions.merge_cells("A10:C10")
        functions["A10"] = "功能信息"
        functions.merge_cells("D10:E10")
        functions["D10"] = "实现信息"
        headers = ["功能ID", "父功能ID", "功能名称", "功能描述", "主要实现模块"]
        for column, value in enumerate(headers, start=1):
            functions.cell(11, column, value)
        for column, value in enumerate(
            ["FUNC-DEMO-01", "", "输入处理", "接收并标准化输入", "MOD-DEMO-A / 输入模块"],
            start=1,
        ):
            functions.cell(12, column, value)
        functions.cell(13, 4, "补充执行格式检查与时间戳校验")
        for column, value in enumerate(headers, start=1):
            functions.cell(14, column, value)
        for column, value in enumerate(
            ["FUNC-DEMO-02", "", "输出处理", "形成输出数据", "MOD-DEMO-B / 输出模块"],
            start=1,
        ):
            functions.cell(15, column, value)

        modules = workbook.create_sheet("模块")
        modules.append(["所属子系统", "模块ID", "缩写", "模块名称", "主要职责"])
        modules.append(["演示子系统 (DSS)", "MOD-DEMO-A", "DMA", "输入模块", "接收输入"])
        modules.append(["演示子系统 (DSS)", "MOD-DEMO-B", "DMB", "输出模块", "形成输出"])

        subsystems = workbook.create_sheet("子系统")
        subsystems.append(["子系统ID", "缩写", "名称", "主要职责", "包含模块"])
        subsystems.append(["SUB-DEMO", "DSS", "演示子系统", "完成演示处理", "DMA（输入模块）、DMB（输出模块）"])

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "function-layout.xlsx"
            workbook.save(path)
            result = extract_function_drafts(str(path))

        self.assertEqual(
            result.stats,
            {
                "systems": 1,
                "subsystems": 1,
                "modules": 2,
                "functions": 2,
                "functions_with_failure_symptoms": 1,
                "interfaces": 1,
                "image_evidence": 0,
                "invalid_records": 0,
            },
        )
        first = next(item for item in result.functions if item.values["function_id"] == "FUNC-DEMO-01")
        self.assertIn("补充执行格式检查", first.values["description"])
        self.assertEqual(first.values["module_id"], "MOD-DEMO-A")
        self.assertEqual(first.values["subsystem_name"], "演示子系统")
        self.assertEqual(first.values["system_name"], "演示综合处理系统")
        self.assertEqual(first.values["failure_symptom_list"], ["输入不更新", "状态字冻结"])
        self.assertGreaterEqual(len(first.source_refs), 3)
        self.assertTrue(first.raw_data["source_records"])
        module_a = next(
            item for item in result.components if item.values["component_id"] == "MOD-DEMO-A"
        )
        self.assertEqual(module_a.values["description"], "接收输入")

        edge = result.interfaces[0]
        self.assertTrue(edge.valid)
        self.assertEqual(edge.values["source_component_id"], "MOD-DEMO-A")
        self.assertEqual(edge.values["target_component_id"], "MOD-DEMO-B")
        self.assertEqual(edge.values["relation_type"], "data_flow")
        self.assertEqual(edge.values["protocol"], "LOCAL-BUS")

    def test_labelled_paragraph_and_explicit_arrow_are_deterministic_sources(self):
        document = Document()
        document.add_paragraph("功能ID：FUNC-TEXT-01")
        document.add_paragraph("功能名称：状态汇聚")
        document.add_paragraph("功能描述：汇聚两路显式状态输入")
        document.add_paragraph("实现模块：MOD-TEXT-A")
        document.add_paragraph("MOD-TEXT-A -> MOD-TEXT-B：健康状态字")

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "labelled.docx"
            document.save(path)
            result = extract_function_drafts(str(path))

        self.assertEqual(len(result.functions), 1)
        self.assertEqual(result.functions[0].values["function_id"], "FUNC-TEXT-01")
        self.assertEqual(result.functions[0].values["module_id"], "MOD-TEXT-A")
        self.assertEqual(len(result.interfaces), 1)
        self.assertTrue(result.interfaces[0].valid)
        self.assertEqual(result.interfaces[0].values["source_component_id"], "MOD-TEXT-A")
        self.assertEqual(result.interfaces[0].values["target_component_id"], "MOD-TEXT-B")
        self.assertIn(
            "interface_endpoint_not_declared",
            [issue.code for issue in result.interfaces[0].issues],
        )

    def test_contained_module_declaration_fills_missing_module_parent(self):
        """A subsystem's explicit 包含模块 field can be the only parent source."""

        workbook = Workbook()
        subsystems = workbook.active
        subsystems.title = "子系统"
        subsystems.append(["子系统ID", "缩写", "名称", "主要职责", "包含模块"])
        subsystems.append(
            [
                "SUB-PWR",
                "PWR",
                "电源与配电子系统",
                "提供受控电源",
                "ICU（输入调理模块）",
            ]
        )
        modules = workbook.create_sheet("模块")
        # The module table deliberately omits 所属子系统.  The parser must
        # retain the explicit structured declaration above rather than leave
        # this module attached only through an incidental merged source row.
        modules.append(["模块ID", "缩写", "模块名称", "主要职责"])
        modules.append(["MOD-PWR-ICU", "ICU", "输入调理模块", "采样和输入调理"])

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "contained-module-parent.xlsx"
            workbook.save(path)
            result = extract_function_drafts(str(path))

        module = next(
            item
            for item in result.components
            if item.values["component_id"] == "MOD-PWR-ICU"
        )
        self.assertEqual("SUB-PWR", module.values["parent_component_id"])
        declaration_sources = [
            source
            for source in module.raw_data["source_records"]
            if source.get("role") == "subsystem_module_declaration"
        ]
        self.assertEqual(1, len(declaration_sources))
        self.assertEqual(
            "SUB-PWR", declaration_sources[0]["raw_values"]["声明父组件ID"]
        )

    def test_contained_module_declaration_keeps_conflicting_existing_parent(self):
        workbook = Workbook()
        subsystems = workbook.active
        subsystems.title = "子系统"
        subsystems.append(["子系统ID", "缩写", "名称", "主要职责", "包含模块"])
        subsystems.append(
            [
                "SUB-PWR",
                "PWR",
                "电源与配电子系统",
                "提供受控电源",
                "ICU（输入调理模块）",
            ]
        )
        subsystems.append(["SUB-OTHER", "OTH", "其他子系统", "其他职责", ""])
        modules = workbook.create_sheet("模块")
        modules.append(["所属子系统ID", "模块ID", "缩写", "模块名称", "主要职责"])
        modules.append(
            ["SUB-OTHER", "MOD-PWR-ICU", "ICU", "输入调理模块", "采样和输入调理"]
        )

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "contained-module-conflict.xlsx"
            workbook.save(path)
            result = extract_function_drafts(str(path))

        module = next(
            item
            for item in result.components
            if item.values["component_id"] == "MOD-PWR-ICU"
        )
        self.assertEqual("SUB-OTHER", module.values["parent_component_id"])
        self.assertIn(
            "conflicting_declared_module_parent",
            [issue.code for issue in module.issues],
        )
        self.assertEqual(
            ["SUB-OTHER", "SUB-PWR"],
            module.raw_data["conflicting_values"]["parent_component_id"],
        )

    def test_contained_module_declaration_never_uses_substring_alias_as_parent(self):
        workbook = Workbook()
        subsystems = workbook.active
        subsystems.title = "子系统"
        subsystems.append(["子系统ID", "缩写", "名称", "主要职责", "包含模块"])
        subsystems.append(
            ["SUB-PWR", "PWR", "电源与配电子系统", "提供受控电源", "输入滤波模块"]
        )
        modules = workbook.create_sheet("模块")
        modules.append(["模块ID", "缩写", "模块名称", "主要职责"])
        modules.append(["MOD-SHORT", "IN", "输入", "另一个短名称模块"])

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "substring-module-parent.xlsx"
            workbook.save(path)
            result = extract_function_drafts(str(path))

        short_module = next(
            item
            for item in result.components
            if item.values["component_id"] == "MOD-SHORT"
        )
        self.assertFalse(short_module.values.get("parent_component_id"))
        declared = [
            item
            for item in result.components
            if item.values.get("name") == "输入滤波模块"
        ]
        self.assertEqual(1, len(declared))
        self.assertEqual("SUB-PWR", declared[0].values["parent_component_id"])

    def test_contained_module_declaration_supports_both_orders_and_whitespace_lists(self):
        workbook = Workbook()
        subsystems = workbook.active
        subsystems.title = "子系统"
        subsystems.append(["子系统ID", "缩写", "名称", "主要职责", "包含模块"])
        subsystems.append(
            [
                "SUB-DEMO",
                "DEMO",
                "演示子系统",
                "验证显式模块声明",
                (
                    "MOD-A（Alpha模块）\n"
                    "Beta模块（MOD-B）\n"
                    "MOD-C Gamma模块 MOD-D Delta模块\n"
                    "Epsilon模块 MOD-E Zeta模块 MOD-F"
                ),
            ]
        )

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "declared-module-formats.xlsx"
            workbook.save(path)
            result = extract_function_drafts(str(path))

        modules = {
            item.values.get("abbreviation"): item
            for item in result.components
            if item.values.get("component_type") == "module"
        }
        self.assertEqual(
            {
                "MOD-A": "Alpha模块",
                "MOD-B": "Beta模块",
                "MOD-C": "Gamma模块",
                "MOD-D": "Delta模块",
                "MOD-E": "Epsilon模块",
                "MOD-F": "Zeta模块",
            },
            {abbreviation: item.values.get("name") for abbreviation, item in modules.items()},
        )
        self.assertTrue(
            all(item.values.get("parent_component_id") == "SUB-DEMO" for item in modules.values())
        )

    def test_system_parent_does_not_match_short_alias_inside_description(self):
        workbook = Workbook()
        systems = workbook.active
        systems.title = "系统"
        systems.append(["系统ID", "系统缩写", "系统名称", "系统说明"])
        systems.append(["SYS-ICP", "ICP", "综合处理系统", "主系统"])
        subsystems = workbook.create_sheet("子系统")
        subsystems.append(["子系统ID", "所属系统", "缩写", "名称", "主要职责"])
        subsystems.append(
            [
                "SUB-OTHER",
                "这是包含 ICP 字样的外部说明，不是系统引用",
                "OTH",
                "其他子系统",
                "其他职责",
            ]
        )

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "system-parent-substring.xlsx"
            workbook.save(path)
            result = extract_function_drafts(str(path))

        subsystem = next(
            item
            for item in result.components
            if item.values.get("component_id") == "SUB-OTHER"
        )
        self.assertFalse(subsystem.values.get("parent_component_id"))
        self.assertIn(
            "unresolved_parent_system", [issue.code for issue in subsystem.issues]
        )

    def test_module_parent_does_not_match_short_alias_inside_description(self):
        workbook = Workbook()
        subsystems = workbook.active
        subsystems.title = "子系统"
        subsystems.append(["子系统ID", "缩写", "名称", "主要职责"])
        subsystems.append(["SUB-PWR", "PWR", "电源子系统", "供电"])
        modules = workbook.create_sheet("模块")
        modules.append(["所属子系统", "模块ID", "缩写", "模块名称", "主要职责"])
        modules.append(
            [
                "用于说明 PWR 边界但不是子系统引用",
                "MOD-OTHER",
                "OTH",
                "其他模块",
                "其他职责",
            ]
        )

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "module-parent-substring.xlsx"
            workbook.save(path)
            result = extract_function_drafts(str(path))

        module = next(
            item
            for item in result.components
            if item.values.get("component_id") == "MOD-OTHER"
        )
        self.assertFalse(module.values.get("parent_component_id"))
        self.assertIn(
            "unresolved_parent_subsystem", [issue.code for issue in module.issues]
        )

    def test_function_module_does_not_match_short_alias_inside_description(self):
        workbook = Workbook()
        modules = workbook.active
        modules.title = "模块"
        modules.append(["模块ID", "缩写", "模块名称", "主要职责"])
        modules.append(["MOD-PWR-ICU", "ICU", "输入调理模块", "输入调理"])
        functions = workbook.create_sheet("功能")
        functions.append(["功能ID", "功能名称", "功能描述", "实现模块"])
        functions.append(
            [
                "FUNC-OTHER",
                "其他功能",
                "不属于输入调理模块",
                "用于监控 ICU 状态的外部说明，不是实现模块引用",
            ]
        )

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "function-module-substring.xlsx"
            workbook.save(path)
            result = extract_function_drafts(str(path))

        function = next(
            item
            for item in result.functions
            if item.values.get("function_id") == "FUNC-OTHER"
        )
        self.assertFalse(function.values.get("module_id"))
        self.assertIn(
            "unresolved_function_module", [issue.code for issue in function.issues]
        )

    def test_interface_endpoint_does_not_match_short_alias_inside_description(self):
        workbook = Workbook()
        modules = workbook.active
        modules.title = "模块"
        modules.append(["模块ID", "缩写", "模块名称", "主要职责"])
        modules.append(["MOD-PWR-ICU", "ICU", "输入调理模块", "输入调理"])
        modules.append(["MOD-OUT", "OUT", "输出模块", "输出处理"])
        interfaces = workbook.create_sheet("接口")
        interfaces.append(["接口ID", "发送方", "接收方", "关系类型"])
        interfaces.append(
            [
                "IF-OTHER",
                "外部描述中提到 ICU，但该文本不是接口端点",
                "MOD-OUT",
                "数据流",
            ]
        )

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "interface-endpoint-substring.xlsx"
            workbook.save(path)
            result = extract_function_drafts(str(path))

        interface = next(
            item
            for item in result.interfaces
            if item.values.get("interface_id") == "IF-OTHER"
        )
        self.assertFalse(interface.values.get("source_component_id"))
        self.assertEqual("MOD-OUT", interface.values.get("target_component_id"))
        self.assertFalse(interface.valid)
        self.assertIn(
            "unresolved_interface_endpoint", [issue.code for issue in interface.issues]
        )

    def test_image_only_document_never_creates_authoritative_rows(self):
        document = Document()
        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "architecture.png"
            Image.new("RGB", (8, 8), "white").save(image_path)
            document.add_picture(str(image_path))
            path = Path(directory) / "image-only.docx"
            document.save(path)
            result = extract_function_drafts(str(path))

        self.assertEqual(result.components, [])
        self.assertEqual(result.functions, [])
        self.assertEqual(result.interfaces, [])
        self.assertEqual(len(result.image_evidence), 1)
        self.assertEqual(result.image_evidence[0]["authority"], "unconfirmed_extraction")
        self.assertFalse(result.image_evidence[0]["eligible_for_authoritative_records"])


@unittest.skipUnless(TEST_DOCS.is_dir(), "test-docs 样例目录不存在")
class FunctionTestDocsAcceptanceTests(unittest.TestCase):
    def test_architecture_sample_has_expected_authoritative_counts_and_merges(self):
        result = extract_function_drafts(
            str(TEST_DOCS / "01_系统架构与功能说明.docx")
        )

        self.assertEqual(result.stats["systems"], 1)
        self.assertEqual(result.stats["subsystems"], 6)
        self.assertEqual(result.stats["modules"], 24)
        self.assertEqual(result.stats["functions"], 30)
        self.assertEqual(result.stats["functions_with_failure_symptoms"], 24)
        self.assertEqual(result.stats["interfaces"], 0)
        self.assertEqual(result.stats["image_evidence"], 4)
        self.assertEqual(result.stats["invalid_records"], 0)

        leaf = next(item for item in result.functions if item.values["function_id"] == "FUNC-110")
        self.assertEqual(leaf.values["module_id"], "MOD-PWR-ICU")
        self.assertEqual(leaf.values["module_name"], "输入调理模块")
        self.assertEqual(leaf.values["subsystem_name"], "电源与配电子系统")
        self.assertEqual(leaf.values["system_name"], "综合信息处理与显示系统")
        self.assertEqual(len(leaf.source_refs), 2)
        self.assertEqual(len(leaf.values["failure_symptom_list"]), 3)
        row = leaf.to_row()
        self.assertTrue(row["raw_data"]["raw_values"])
        self.assertTrue(row["source_ref"])
        self.assertIn("normalization_warnings", row)

        # SUB-PWR's 包含模块 declarations are the explicit parent source for
        # module rows that do not repeat 所属子系统 in the module table.
        input_module = next(
            item
            for item in result.components
            if item.values["component_id"] == "MOD-PWR-ICU"
        )
        self.assertEqual("SUB-PWR", input_module.values["parent_component_id"])

        # The sample has explicit data-flow prose and four diagrams, but no
        # authoritative interface table or arrow expression. Neither is guessed.
        self.assertEqual(result.interfaces, [])
        self.assertTrue(
            all(
                evidence["eligible_for_authoritative_records"] is False
                for evidence in result.image_evidence
            )
        )


if __name__ == "__main__":
    unittest.main()
