import base64
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from backend.core import docx_parser as docx_parser_module
from backend.core.parsers.base import ParseContext
from backend.core.parsers.diagram import DiagramParser
from backend.core.parsers.diagram_extraction import extract_fta_from_ir, extract_rbd_from_ir
from backend.core.parsers.document_ir import load_document_ir
from backend.core.parsers.excel_row import extract_domain_drafts
from backend.core.parsers.fault_history import extract_history_drafts
from backend.core.parsers.function_extraction import extract_function_drafts


_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


class ParserSourceNameStabilityTests(unittest.TestCase):
    @staticmethod
    def _save_docx_twice(document: Document, directory: str, stem: str):
        first = Path(directory) / f"materialized-{stem}-aaa.docx"
        second = Path(directory) / f"materialized-{stem}-bbb.docx"
        document.save(first)
        document.save(second)
        return first, second

    def test_document_ir_source_override_is_basename_only_for_docx_and_xlsx(self):
        document = Document()
        document.add_paragraph("稳定来源")
        workbook = Workbook()
        workbook.active.append(["字段", "值"])

        with tempfile.TemporaryDirectory() as directory:
            docx_path = Path(directory) / "materialized-random.docx"
            xlsx_path = Path(directory) / "materialized-random.xlsx"
            document.save(docx_path)
            workbook.save(xlsx_path)
            docx_ir = load_document_ir(
                str(docx_path), source_name=r"C:\incoming\架构说明.docx"
            )
            xlsx_ir = load_document_ir(
                str(xlsx_path), source_name="/incoming/HMC清单.xlsx"
            )

        self.assertEqual(docx_ir.file_name, "架构说明.docx")
        self.assertEqual(str(docx_ir.paragraphs[0].source).split("#", 1)[0], "架构说明.docx")
        self.assertEqual(xlsx_ir.file_name, "HMC清单.xlsx")
        self.assertEqual(str(xlsx_ir.tables[0].source).split("#", 1)[0], "HMC清单.xlsx")

    def test_hmc_variant_and_fmeca_fallback_are_stable_across_materializations(self):
        hmc = Document()
        hmc_table = hmc.add_table(rows=2, cols=5)
        for column, value in enumerate(("LCN", "HAMC", "HSC", "故障现象", "维修方法")):
            hmc_table.cell(0, column).text = value
        for column, value in enumerate(("00123456", "**", "0001", "显示黑屏", "检查显示链路")):
            hmc_table.cell(1, column).text = value

        fmeca = Document()
        fmeca_table = fmeca.add_table(rows=2, cols=3)
        for column, value in enumerate(("序号", "故障模式", "故障原因")):
            fmeca_table.cell(0, column).text = value
        for column, value in enumerate(("1", "处理器挂起", "软件死锁")):
            fmeca_table.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            hmc_paths = self._save_docx_twice(hmc, directory, "hmc")
            fmeca_paths = self._save_docx_twice(fmeca, directory, "fmeca")
            hmc_drafts = [
                extract_domain_drafts(
                    str(path), "hmc", source_name="02_HMC故障代码手册.docx"
                )[0]
                for path in hmc_paths
            ]
            fmeca_drafts = [
                extract_domain_drafts(
                    str(path), "fmeca", source_name="03_FMECA分析报告.docx"
                )[0]
                for path in fmeca_paths
            ]

        self.assertEqual(hmc_drafts[0].source_ref, hmc_drafts[1].source_ref)
        self.assertEqual(
            hmc_drafts[0].values["variant_id"], hmc_drafts[1].values["variant_id"]
        )
        self.assertNotIn("materialized-", hmc_drafts[0].source_ref)
        self.assertEqual(fmeca_drafts[0].source_ref, fmeca_drafts[1].source_ref)
        self.assertEqual(fmeca_drafts[0].values["code"], fmeca_drafts[1].values["code"])
        self.assertTrue(fmeca_drafts[0].values["code"].startswith("FMECA-SRC-"))

    def test_function_interface_and_history_evidence_are_stable(self):
        architecture = Document()
        architecture.add_paragraph("功能ID：FUNC-TEXT-01")
        architecture.add_paragraph("功能名称：状态汇聚")
        architecture.add_paragraph("实现模块：MOD-TEXT-A")
        architecture.add_paragraph("MOD-TEXT-A -> MOD-TEXT-B：健康状态字")

        history = Document()
        history.add_heading("2.1 HIS-001 显示异常排故记录", level=2)
        history.add_paragraph("故障现象：主画面黑屏")
        history.add_paragraph("1. 查询显示任务心跳")

        with tempfile.TemporaryDirectory() as directory:
            architecture_paths = self._save_docx_twice(architecture, directory, "function")
            history_paths = self._save_docx_twice(history, directory, "history")
            function_results = [
                extract_function_drafts(
                    str(path), source_name="01_系统架构与功能说明.docx"
                )
                for path in architecture_paths
            ]
            history_results = [
                extract_history_drafts(
                    str(path), source_name="07_故障历史.docx"
                )
                for path in history_paths
            ]

        first_interface = function_results[0].interfaces[0]
        second_interface = function_results[1].interfaces[0]
        self.assertEqual(first_interface.record_key, second_interface.record_key)
        self.assertEqual(first_interface.source_ref, second_interface.source_ref)
        self.assertNotIn("materialized-", first_interface.source_ref)
        self.assertEqual(history_results[0][0].source_ref, history_results[1][0].source_ref)
        self.assertEqual(
            history_results[0][0].values["step_records"],
            history_results[1][0].values["step_records"],
        )

    def test_fta_and_rbd_source_dependent_candidate_hashes_are_stable(self):
        fta = Document()
        fta.add_heading("2.1 显示链路 FTA", level=1)
        metadata = fta.add_table(rows=2, cols=2)
        metadata.cell(0, 0).text, metadata.cell(0, 1).text = "故障树编号", "FTA-REAL-01"
        metadata.cell(1, 0).text, metadata.cell(1, 1).text = "顶事件", "显示链路失效"
        events = fta.add_table(rows=2, cols=3)
        for column, value in enumerate(("事件ID", "事件名称", "事件概率")):
            events.cell(0, column).text = value
        for column, value in enumerate(("E-01", "处理模块失效", "0.001")):
            events.cell(1, column).text = value
        cuts = fta.add_table(rows=2, cols=2)
        cuts.cell(0, 0).text, cuts.cell(0, 1).text = "序号", "最小割集"
        cuts.cell(1, 0).text, cuts.cell(1, 1).text = "1", "处理模块失效"

        rbd = Document()
        rbd.add_heading("4.1 参数证据 RBD", level=1)
        parameters = rbd.add_table(rows=2, cols=3)
        for column, value in enumerate(("模块ID", "模块名称", "失效率λ（1/h）")):
            parameters.cell(0, column).text = value
        for column, value in enumerate(("MOD-A", "处理模块", "1.2E-5")):
            parameters.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            fta_paths = self._save_docx_twice(fta, directory, "fta")
            rbd_paths = self._save_docx_twice(rbd, directory, "rbd")
            fta_candidates = [
                extract_fta_from_ir(
                    load_document_ir(str(path), source_name="04_FTA故障树分析报告.docx")
                )[0].candidate
                for path in fta_paths
            ]
            rbd_candidates = [
                extract_rbd_from_ir(
                    load_document_ir(str(path), source_name="05_RBD可靠性分析报告.docx")
                )[0].candidate
                for path in rbd_paths
            ]

        for candidates in (fta_candidates, rbd_candidates):
            self.assertEqual(candidates[0]["candidate_id"], candidates[1]["candidate_id"])
            self.assertEqual(
                candidates[0]["source_fingerprint"], candidates[1]["source_fingerprint"]
            )
            self.assertEqual(candidates[0]["source_locator"], candidates[1]["source_locator"])
            self.assertNotIn("materialized-", candidates[0]["source_locator"])

    def test_docx_and_direct_image_evidence_use_the_original_name(self):
        document = Document()
        document.add_heading("1 系统架构", level=1)

        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "embedded.png"
            image_path.write_bytes(_PNG_1X1)
            document.add_picture(str(image_path))
            docx_paths = self._save_docx_twice(document, directory, "images")
            with patch.object(docx_parser_module.artifact_store, "upload_file"):
                parsed = [
                    docx_parser_module.DocumentParser(
                        str(path),
                        "doc-images",
                        source_name=r"C:\incoming\01_系统架构与功能说明.docx",
                    ).process()
                    for path in docx_paths
                ]

            direct_path = Path(directory) / "materialized-direct-random.png"
            direct_path.write_bytes(_PNG_1X1)
            ctx = ParseContext(
                doc_id="doc-rbd",
                doc_type="rbd",
                file_name=r"C:\incoming\05_RBD可靠性图.png",
            )
            parser = DiagramParser()
            with patch("backend.core.parsers.diagram.artifact_store.upload_file") as upload:
                direct_blocks = parser._load_blocks(str(direct_path), ctx)

        image_refs = [
            next(
                image["source_ref"]
                for block in item["blocks"].values()
                for image in block["images"]
            )
            for item in parsed
        ]
        self.assertEqual(image_refs[0], image_refs[1])
        self.assertTrue(image_refs[0].startswith("01_系统架构与功能说明.docx#"))
        self.assertNotIn("materialized-", image_refs[0])
        block = direct_blocks[0][1]
        self.assertEqual(block["title"], "05_RBD可靠性图")
        self.assertEqual(
            block["images"][0]["source_ref"], "05_RBD可靠性图.png#direct-image"
        )
        self.assertEqual(upload.call_args.kwargs["original_name"], "05_RBD可靠性图.png")


if __name__ == "__main__":
    unittest.main()
