import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from backend.core.parsers.excel_row import extract_domain_drafts
from backend.core.parsers.records import RecordDraft, SourceRef


class TableAdapterTests(unittest.TestCase):
    def _save(self, document: Document, directory: str, name: str) -> str:
        path = Path(directory) / name
        document.save(path)
        return str(path)

    def test_hmc_parses_split_header_preserves_raw_and_pads_display_hsc(self):
        document = Document()
        table = document.add_table(rows=3, cols=5)
        table.cell(0, 0).merge(table.cell(0, 2)).text = "HMC代码"
        table.cell(0, 3).text = "故障描述现象"
        table.cell(0, 4).text = "故障分析程序/维修方法"
        for column, value in enumerate(("LCN", "HAMC", "HSC", "故障描述现象", "故障分析程序/维修方法")):
            table.cell(1, column).text = value
        for column, value in enumerate(("00123456", "**", "1", "显示黑屏", "检查显示链路")):
            table.cell(2, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = self._save(document, directory, "hmc.docx")
            drafts = extract_domain_drafts(path, "hmc")

        self.assertEqual(len(drafts), 1)
        self.assertTrue(drafts[0].valid)
        self.assertEqual(drafts[0].values["lcn_raw"], "00123456")
        self.assertEqual(drafts[0].values["hsc_raw"], "1")
        self.assertEqual(drafts[0].values["hsc"], "0001")
        self.assertEqual(drafts[0].values["hmc_code"], "00123456**0001")
        self.assertIn("row=3", drafts[0].source_ref)

    def test_fmeca_prefers_business_code_and_does_not_invent_numeric_probability(self):
        document = Document()
        table = document.add_table(rows=2, cols=6)
        headers = ["序号", "代码", "故障模式", "故障原因", "严酷度", "故障概率"]
        values = ["1", "FM-REAL-001", "处理器挂起", "软件死锁", "Ⅱ类", "C"]
        for column, value in enumerate(headers):
            table.cell(0, column).text = value
        for column, value in enumerate(values):
            table.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = self._save(document, directory, "fmeca.docx")
            drafts = extract_domain_drafts(path, "fmeca")

        self.assertEqual(len(drafts), 1)
        row = drafts[0].values
        self.assertEqual(row["code"], "FM-REAL-001")
        self.assertEqual(row["sequence_no"], "1")
        self.assertEqual(row["severity"], "II")
        self.assertIsNone(row["probability"])
        self.assertEqual(row["probability_class"], "C")
        self.assertEqual(row["probability_level"], "C")
        self.assertIsNone(row["probability_value"])
        self.assertIn("probability_category_preserved", [issue.code for issue in drafts[0].issues])

    def test_mfl_preserves_status_enum_unknown_and_repair_action(self):
        document = Document()
        table = document.add_table(rows=2, cols=6)
        headers = [
            "MFL", "模块故障模式", "API是否可直接上报", "API名称",
            "是否进行平台级摸底实验", "维修处置",
        ]
        values = ["MFL-001", "模块失效", "是", "GetStatus", "待验证", "更换模块并复测"]
        for column, value in enumerate(headers):
            table.cell(0, column).text = value
        for column, value in enumerate(values):
            table.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = self._save(document, directory, "mfl.docx")
            drafts = extract_domain_drafts(path, "mfl")

        self.assertEqual(len(drafts), 1)
        row = drafts[0].values
        self.assertTrue(row["api_reportable"])
        self.assertIsNone(row["platform_test"])
        self.assertEqual(row["platform_test_status"], "待验证")
        self.assertEqual(row["repair_action"], "更换模块并复测")
        self.assertEqual(row["maintenance_action"], "更换模块并复测")

    def test_fmeca_keeps_numeric_probability_separate_from_display_text(self):
        document = Document()
        table = document.add_table(rows=2, cols=4)
        for column, value in enumerate(("代码", "故障模式", "故障原因", "故障概率")):
            table.cell(0, column).text = value
        for column, value in enumerate(("FM-NUM-001", "电源中断", "供电失效", "2.5%")):
            table.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = self._save(document, directory, "fmeca-numeric.docx")
            drafts = extract_domain_drafts(path, "fmeca")

        row = drafts[0].values
        self.assertEqual(row["probability_raw"], "2.5%")
        self.assertEqual(row["probability_value"], "0.025")
        self.assertEqual(row["probability"], "0.025")
        self.assertEqual(row["probability_level"], "")

    def test_hmc_vertical_merge_creates_source_variants_not_overwrites(self):
        document = Document()
        table = document.add_table(rows=3, cols=5)
        for column, value in enumerate(("LCN", "HAMC", "HSC", "故障现象", "维修方法")):
            table.cell(0, column).text = value
        for column, value in enumerate(("00123456", "**", "0001")):
            table.cell(1, column).merge(table.cell(2, column)).text = value
        table.cell(1, 3).text, table.cell(1, 4).text = "黑屏", "检查供电"
        table.cell(2, 3).text, table.cell(2, 4).text = "无字符", "检查视频链路"

        with tempfile.TemporaryDirectory() as directory:
            path = self._save(document, directory, "hmc-variants.docx")
            drafts = extract_domain_drafts(path, "hmc")

        self.assertEqual(len(drafts), 2)
        self.assertTrue(all(item.valid for item in drafts))
        self.assertEqual({item.values["hmc_code"] for item in drafts}, {"00123456**0001"})
        self.assertEqual(len({item.values["variant_id"] for item in drafts}), 2)
        self.assertEqual({item.values["fault_desc"] for item in drafts}, {"黑屏", "无字符"})
        self.assertIn(
            "merged_cell_inherited",
            [issue.code for issue in drafts[1].issues],
        )

    def test_hmc_complete_code_conflict_fails_closed_and_preserves_raw_values(self):
        document = Document()
        table = document.add_table(rows=2, cols=6)
        headers = ["HMC代码", "LCN", "HAMC", "HSC", "故障现象", "维修方法"]
        values = ["00123456**0002", "00123456", "**", "0001", "黑屏", "检查供电"]
        for column, value in enumerate(headers):
            table.cell(0, column).text = value
        for column, value in enumerate(values):
            table.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = self._save(document, directory, "hmc-conflict.docx")
            draft = extract_domain_drafts(path, "hmc")[0]

        self.assertFalse(draft.valid)
        self.assertEqual(draft.values["hmc_code_raw"], "00123456**0002")
        self.assertEqual(draft.values["hsc_raw"], "0001")
        self.assertIn("inconsistent_hmc_parts", [issue.code for issue in draft.errors])

    def test_hmc_normalizes_fullwidth_code_glyphs_but_keeps_raw_values(self):
        document = Document()
        table = document.add_table(rows=2, cols=5)
        headers = ["LCN", "HAMC", "HSC", "故障现象", "维修方法"]
        values = ["００１２３４５６", "＊＊", "１", "黑屏", "检查供电"]
        for column, value in enumerate(headers):
            table.cell(0, column).text = value
        for column, value in enumerate(values):
            table.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = self._save(document, directory, "hmc-fullwidth.docx")
            draft = extract_domain_drafts(path, "hmc")[0]

        self.assertTrue(draft.valid)
        self.assertEqual(draft.values["lcn_raw"], "００１２３４５６")
        self.assertEqual(draft.values["hamc_raw"], "＊＊")
        self.assertEqual(draft.values["hsc_raw"], "１")
        self.assertEqual(draft.values["hmc_code"], "00123456**0001")

    def test_fmeca_scientific_probability_and_severity_contract(self):
        document = Document()
        table = document.add_table(rows=2, cols=5)
        headers = ["代码", "故障模式", "故障原因", "严酷度", "故障概率"]
        values = ["FM-SCI-01", "丢帧", "时钟漂移", "Ⅱ类", "1.25E-12"]
        for column, value in enumerate(headers):
            table.cell(0, column).text = value
        for column, value in enumerate(values):
            table.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = self._save(document, directory, "fmeca-science.docx")
            row = extract_domain_drafts(path, "fmeca")[0].values

        self.assertEqual(row["severity_raw"], "Ⅱ类")
        self.assertEqual(row["severity_normalized"], "II")
        self.assertEqual(row["probability_raw"], "1.25E-12")
        self.assertEqual(row["probability_value"], "1.25E-12")

    def test_mfl_partial_completion_is_status_not_true(self):
        document = Document()
        table = document.add_table(rows=2, cols=3)
        headers = ["MFL", "模块故障模式", "平台级摸底实验状态"]
        values = ["MFL-STATUS-01", "处理器异常", "部分完成"]
        for column, value in enumerate(headers):
            table.cell(0, column).text = value
        for column, value in enumerate(values):
            table.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = self._save(document, directory, "mfl-status.docx")
            row = extract_domain_drafts(path, "mfl")[0].values

        self.assertIsNone(row["platform_test"])
        self.assertEqual(row["platform_test_status"], "部分完成")

    def test_legacy_xls_numeric_code_is_explicitly_rejected(self):
        draft = RecordDraft(
            record_type="hmc",
            values={
                "hmc_code": "",
                "lcn": "123456",
                "hamc": "**",
                "hsc": "1",
                "fault_desc": "显示异常",
                "repair_method": "检查显示模块",
            },
            raw_data={
                "field_cell_metadata": {
                    "lcn": {"raw_value": 123456, "raw_type": "int"},
                    "hamc": {"raw_value": "**", "raw_type": "str"},
                    "hsc": {"raw_value": 1, "raw_type": "int"},
                }
            },
            source_refs=[SourceRef(file_name="legacy.xls", element_type="row", row_index=2)],
        )
        with patch(
            "backend.core.parsers.excel_row.engine_extract_table_drafts",
            return_value=[draft],
        ):
            result = extract_domain_drafts("legacy.xls", "hmc")[0]

        self.assertFalse(result.valid)
        codes = [issue.code for issue in result.issues]
        self.assertIn("legacy_xls_format_limit", codes)
        self.assertIn("legacy_xls_numeric_code_unsafe", codes)


if __name__ == "__main__":
    unittest.main()
