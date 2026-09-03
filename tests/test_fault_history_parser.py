import tempfile
import unittest
from pathlib import Path

from docx import Document

from backend.core.parsers.fault_history import extract_history_drafts


class FaultHistoryExtractionTests(unittest.TestCase):
    def test_combines_heading_key_value_table_and_multiple_step_formats(self):
        document = Document()
        document.add_heading("2.1 HIS-REAL-001 显示异常排故记录", level=2)
        table = document.add_table(rows=4, cols=2)
        pairs = [
            ("故障编号", "HIS-REAL-001"),
            ("故障现象", "主画面黑屏但后台任务正常"),
            ("关联系统设计模块", "DSP / DPM"),
            ("初始HMC代码", "91840001-**-0001"),
        ]
        for row, pair in zip(table.rows, pairs):
            row.cells[0].text, row.cells[1].text = pair
        document.add_paragraph("使用的排故方法：状态查询 + 替换验证")
        document.add_paragraph("处理步骤（含API调用）：")
        document.add_paragraph("1. 查询显示任务心跳")
        document.add_paragraph("第2步：切换备用显示通道")
        document.add_paragraph("（3）更换显示处理模块并复测")
        document.add_paragraph("故障原因分析：显示任务死锁")
        document.add_paragraph("最终解决方案：升级软件并更换异常模块")
        document.add_paragraph(
            "关联文档索引：FMECA→FM-REAL-001；FTA→FTA-REAL-001/EVT-REAL-01；"
            "RBD→RBD-REAL-001；MFL→MFL-REAL-001"
        )

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "history.docx"
            document.save(path)
            drafts = extract_history_drafts(str(path))

        self.assertEqual(len(drafts), 1)
        draft = drafts[0]
        self.assertTrue(draft.valid)
        self.assertEqual(draft.values["case_id"], "HIS-REAL-001")
        self.assertEqual(len(draft.values["steps"]), 3)
        self.assertEqual(draft.values["hmc_code"], "91840001**0001")
        self.assertEqual(draft.values["fmeca_code"], "FM-REAL-001")
        self.assertEqual(draft.values["fta_ref"], "FTA-REAL-001")
        self.assertEqual(draft.values["rbd_ref"], "RBD-REAL-001")
        self.assertEqual(draft.values["mfl_code"], "MFL-REAL-001")
        self.assertEqual(draft.values["document_links"]["rbd_refs"], ["RBD-REAL-001"])
        self.assertEqual(len(draft.values["step_records"]), 3)
        self.assertTrue(all(item["source_ref"] for item in draft.values["step_records"]))

    def test_uses_wide_history_table_when_no_detailed_cases_exist(self):
        document = Document()
        table = document.add_table(rows=2, cols=5)
        headers = ["故障编号", "故障现象", "根本原因", "最终方案", "预防措施"]
        values = ["CASE-2026-001", "接口间歇中断", "连接器松动", "重新固定连接器", "增加力矩检查"]
        for column, value in enumerate(headers):
            table.cell(0, column).text = value
        for column, value in enumerate(values):
            table.cell(1, column).text = value

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "history-table.docx"
            document.save(path)
            drafts = extract_history_drafts(str(path))

        self.assertEqual(len(drafts), 1)
        self.assertTrue(drafts[0].valid)
        self.assertEqual(drafts[0].values["case_id"], "CASE-2026-001")
        self.assertEqual(drafts[0].values["root_cause"], "连接器松动")


if __name__ == "__main__":
    unittest.main()
